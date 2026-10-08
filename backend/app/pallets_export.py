"""Eksport wywołania palet do xlsx (openpyxl) — załącznik maila do DLT.

Wywołanie z autami (PalletCallTruck) -> arkusz per auto (miejsce, materiał, HU,
ilość palet); pozycje bez przydziału trafiają do arkusza „Bez auta".
Wywołanie bez aut -> jeden arkusz jak dotychczas.
"""
from __future__ import annotations

import io

from openpyxl import Workbook

from .exports import append_row

_COLS = [
    ("produkt", "Produkt"), ("krotki_opis", "Opis"),
    ("ilosc_pal", "Ilość palet"), ("data_dostawy", "Data dostawy"), ("note", "Uwagi"),
]
_TRUCK_HEADER = ["Miejsce", "Produkt", "Opis", "HU", "Palety", "Miejsca pal.",
                 "Data dostawy", "Uwagi"]


def _places(line) -> int:
    import math
    return math.ceil(float(line.pallets)) if line.pallets else int(float(line.ilosc_pal))


def _truck_sheet(ws, call, lines) -> None:
    append_row(ws, ["Wywołanie", call.number])
    append_row(ws, _TRUCK_HEADER)
    slot = 1
    for line in lines:
        append_row(ws, [slot, line.produkt, line.krotki_opis, line.hu_numbers,
                        float(line.pallets) if line.pallets is not None else None,
                        _places(line), line.data_dostawy, line.note])
        slot += _places(line)


def build_xlsx(call) -> bytes:
    wb = Workbook()
    trucks = getattr(call, "trucks", None) or []
    if trucks:
        wb.remove(wb.active)
        by_truck = {t.id: [] for t in trucks}
        loose = []
        for line in call.lines:
            by_truck.get(line.truck_id, loose).append(line)
        for t in sorted(trucks, key=lambda t: t.ordinal):
            _truck_sheet(wb.create_sheet(f"Auto {t.ordinal}"), call, by_truck[t.id])
        if loose:
            _truck_sheet(wb.create_sheet("Bez auta"), call, loose)
    else:
        ws = wb.active
        ws.title = "Wywolanie"
        append_row(ws, ["Wywołanie", call.number])
        append_row(ws, [label for _, label in _COLS])
        for line in call.lines:
            append_row(ws, [getattr(line, attr) for attr, _ in _COLS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_filename(call) -> str:
    return f"wywolanie_{call.number}.xlsx"
