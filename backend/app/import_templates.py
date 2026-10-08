"""Wzory plików importu (xlsx do pobrania przy każdym imporcie, 2026-09-24).

Arkusz „Dane" = nagłówki w brzmieniu, które rozpoznaje parser danego importu + jeden
przykładowy wiersz (daty jako prawdziwe komórki daty — parsery SAP nie czytają dat-tekstów).
Arkusz „Instrukcja" = opis kolumn. „Dane" jest PIERWSZY — większość parserów czyta
worksheets[0]. Test (tests/test_import_templates.py) przepuszcza każdy wzór przez
prawdziwy parser, więc zmiana nagłówków w parserze bez zmiany wzoru = czerwony test.
"""
import datetime
import io

from openpyxl import Workbook
from openpyxl.styles import Font

from .exports import append_row

D = datetime.date(2026, 10, 15)

# kind -> (tytuł, [(nagłówek, przykład, opis kolumny)]); „*" w opisie = kolumna wymagana
TEMPLATES: dict[str, tuple[str, list[tuple[str, object, str]]]] = {
    "materials": ("Materiały (master data)", [
        ("ref_code", "REF12345", "* numer materiału (REF) — klucz; ponowny import aktualizuje"),
        ("opis_pl", "Krzesło ogrodowe", "nazwa PL"),
        ("opis_en", "Garden chair", "nazwa EN"),
        ("rodzina", "Meble", "rodzina / grupa"),
        ("jm", "SZT", "podstawowa jednostka miary"),
        ("tariff_cn", "94017100", "kod CN (taryfa)"),
        ("customs_code", "9401710000", "kod celny"),
        ("vat", "23", "stawka VAT (23, 8, zw, np)"),
        ("sent", "nie", "podlega SENT: tak/nie"),
        ("supplier_codes", "100200", "kody dostawców (po przecinku)"),
        ("producer_code", "GC-01", "kod producenta"),
        ("ean", "5901234123457", "EAN sztuki (8/13/14 cyfr)"),
        ("kar_ean", "15901234123454", "EAN kartonu"),
        ("kar_wymiar", "60x40x40", "wymiar kartonu"),
    ]),
    "marm": ("MARM — jednostki materiałów (SAP)", [
        ("MATNR", "REF12345", "* materiał"),
        ("MEINH", "KAR", "* jednostka alternatywna (SZT, KAR, PAL…)"),
        ("UMREZ", 12, "licznik przelicznika (ile jednostek bazowych)"),
        ("UMREN", 1, "mianownik przelicznika"),
        ("BRGEW", 14.5, "waga brutto"),
        ("GEWEI", "KG", "jednostka wagi"),
        ("LAENG", 60, "długość"),
        ("BREIT", 40, "szerokość"),
        ("HOEH", 40, "wysokość"),
        ("MEABM", "CM", "jednostka wymiarów"),
        ("VOLUM", 0.096, "objętość"),
        ("VOLEH", "M3", "jednostka objętości"),
    ]),
    "lfa1": ("LFA1 — dostawcy (SAP)", [
        ("Dostawca", "100200", "* numer dostawcy SAP"),
        ("Kraj", "CN", "kod kraju (2 litery)"),
        ("Nazwa 1", "Ningbo Garden Co.", "* nazwa"),
        ("Nazwa 2", "Ltd.", "ciąg dalszy nazwy"),
        ("Miasto", "Ningbo", ""),
        ("Kod poczt.", "315000", ""),
        ("Ulica", "Harbour Rd 1", ""),
        ("Nr ident. VAT", "CN913302", ""),
        ("Centralna blokada księgowania", "", "X = zablokowany w SAP"),
    ]),
    "ekko": ("EKKO — nagłówki zamówień (SAP)", [
        ("Dok.zaopatrz.", "4500617421", "* numer zamówienia — klucz"),
        ("Rodzaj dok.", "NB", ""),
        ("Dostawca", "100200", "* numer dostawcy SAP"),
        ("Utworzone przez", "JKOWALSKI", "kupiec"),
        ("Waluta", "USD", ""),
        ("Wart.całk.", 12500.5, "wartość zamówienia"),
        ("Incoterms", "FOB", ""),
        ("Data dokumentu", datetime.date(2026, 8, 1), "komórka typu DATA"),
        ("Data dostawy", D, "komórka typu DATA"),
        ("Planowana data wysyłki", datetime.date(2026, 9, 10), "komórka typu DATA"),
        ("ASAP", "", "X = pilne"),
    ]),
    "ref": ("REF — pozycje zamówień (SAP)", [
        ("Dok.zaopatrz.", "4500617421", "* numer zamówienia"),
        ("Pozycja", "10", "pozycja (pusta = numer kolejny)"),
        ("Krótki tekst", "Krzesło ogrodowe", ""),
        ("Materiał", "REF12345", "* materiał"),
        ("Ilość zam.", 1200, ""),
        ("Jedn. miary", "SZT", ""),
        ("Waga netto", 900, ""),
        ("Waga brutto", 1100, ""),
        ("Objętość", 9.6, ""),
        ("Jednostka obj.", "M3", ""),
        ("Planowana data wysyłki", datetime.date(2026, 9, 10), "komórka typu DATA"),
    ]),
    "etd": ("Zamówienia ETD (arkusz zakupów)", [
        ("ORDER NO", "4500617421", "* numer zamówienia — klucz"),
        ("ETD", datetime.date(2026, 9, 10), "komórka typu DATA"),
        ("SUPPLIER", "Ningbo Garden Co.", ""),
        ("PI NO", "PI-2026-77", ""),
        ("PRODUCT", "Krzesła ogrodowe", ""),
        ("CBM", 58.5, ""),
        ("PORT OF DEPARTURE", "Ningbo", ""),
        ("CONTAINER", "40HC", ""),
        ("AMOUNT", 12500.5, ""),
        ("FORWARDER", "SPEDALFA", ""),
    ]),
    "queue": ("Kolejka kontenerów (kolejka.xlsx)", [
        ("Dostawca", "Ningbo Garden Co.", "tworzony automatycznie, gdy nowy"),
        ("Statek", "MV DEMO ATLAS", ""),
        ("ETA", D, "komórka typu DATA"),
        ("Nr kontenera", "MSDU0806613", "* numer ISO 6346 — klucz"),
        ("Numer zamówienia", "4500617421", "kilka numerów: w osobnych liniach lub po przecinku"),
        ("Magazyn", "DLT", "tworzony automatycznie, gdy nowy"),
        ("Numer dostawy", "18001", ""),
        ("Spedytor", "SPEDALFA", "tworzony automatycznie, gdy nowy"),
        ("Agencja", "", "tylko z rejestru agencji celnych"),
        ("Status odprawy", "", ""),
        ("Data rozładunku", datetime.date(2026, 10, 20), "data awizacji — komórka typu DATA"),
        ("SENT wymagany", "", ""),
        ("Numer SENT", "", ""),
    ]),
    "dlt": ("Stany DLT (snapshot)", [
        ("Produkt", "REF12345", "* materiał"),
        ("Nr HU", "00159012340000000017", ""),
        ("Ilość", 240, "* ilość"),
        ("Lokalizacja", "A-01-02", ""),
    ]),
    "ports": ("Porty kontenerowe", [
        ("Code", "CNNGB", "* kod UN/LOCODE — klucz"),
        ("Port name", "Ningbo", "* nazwa"),
        ("Country", "CN", ""),
        ("Country name", "China", ""),
    ]),
    "po": ("PO → kontenery (dopisanie numerów)", [
        ("container_no", "MSDU0806613", "* istniejący kontener"),
        ("order_numbers", "4500617421, 4500617412", "* numery do dopisania"),
    ]),
    "paz": ("PAZ — sztuk na paletę", [
        ("produkt", "REF12345", "* materiał — klucz"),
        ("sztuk_na_palete", 240, "* liczba > 0"),
    ]),
    "issues": ("Wydania materiałów (analityka)", [
        ("data", "2026-09-01", "* data jako tekst RRRR-MM-DD"),
        ("produkt", "REF12345", "*"),
        ("ilosc", 24, "*"),
        ("firma", "ACME", ""),
        ("magazyn", "DLT", ""),
        ("kontrahent", "Klient ABC", ""),
    ]),
}


def build_template(kind: str) -> bytes:
    title, cols = TEMPLATES[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = "Dane"
    append_row(ws, [c[0] for c in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    append_row(ws, [c[1] for c in cols])
    for cell in ws[2]:
        if isinstance(cell.value, datetime.date):
            cell.number_format = "DD.MM.YYYY"
    for i, (header, _, _) in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(header) + 4)
    info = wb.create_sheet("Instrukcja")
    append_row(info, [title])
    info["A1"].font = Font(bold=True)
    append_row(info, ["Wiersz 2 w arkuszu „Dane” to przykład — nadpisz go swoimi danymi. "
                      "Kolumny z * są wymagane; pozostałe można usunąć lub zostawić puste."])
    append_row(info, [])
    append_row(info, ["Kolumna", "Opis"])
    for header, _, note in cols:
        append_row(info, [header, note])
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 80
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
