"""Raport Excel z podglądów WinSAD: arkusze Zgłoszenia, Pozycje, Walidacja, Dokumenty.

Kontrakt typów komórek (odbiorca filtruje, sumuje i sprawdza w Excelu):
- kwoty jako liczby — `Decimal` prosto do openpyxl, bez float po drodze; format `#,##0.00`;
  kurs `0.0000`; stawki jako ułamek (12 % → 0,12) w formacie `0.0%`;
- daty jako date/datetime z formatem daty;
- identyfikatory (nr SAD, LRN, CN, TARIC, NIP, EORI, kontener, kody) jako TEKST — zera
  wiodące zostają (TARIC „00”);
- kolumny kontrolne w „Pozycje” to FORMUŁY z komórek tego samego wiersza. Plik trzyma
  angielskie nazwy funkcji i przecinki — polski Excel sam pokaże je ze średnikami.
Wiersze idą przez `exports.append_row`: tekst z PDF zaczynający się od „=” zostaje tekstem.
Danych osobowych nie ma w modelu (parser ich nie czyta), więc nie ma ich i w raporcie."""
import io
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ..exports import append_row
from .parser import Pozycja, Zgloszenie
from .validator import ERROR, OK, WARN, Wynik, najgorszy

MONEY = "#,##0.00"
KURS = "0.0000"
PCT = "0.0%"
MASA = "#,##0.00"
DATA = "yyyy-mm-dd"
DATA_CZAS = "yyyy-mm-dd hh:mm"

FONT = Font(name="Arial", size=10)
FONT_NAGLOWEK = Font(name="Arial", size=10, bold=True)
# kolory jak style „Dobry / Neutralny / Zły” w Excelu
KOLOR_OK, KOLOR_WARN, KOLOR_ERROR = "C6EFCE", "FFEB9C", "FFC7CE"
FILLS = {status: PatternFill(fill_type="solid", start_color=kolor, end_color=kolor)
         for status, kolor in ((OK, KOLOR_OK), (WARN, KOLOR_WARN), (ERROR, KOLOR_ERROR))}


class _Formula(str):
    """Szablon formuły: `{klucz}` → adres komórki kolumny o tym kluczu w bieżącym wierszu."""


@dataclass(frozen=True)
class Kol:
    naglowek: str
    wartosc: Callable[..., object]
    fmt: str | None = None      # number_format; None = ogólny (tekst, liczby całkowite)
    szer: float = 14
    klucz: str = ""             # nazwa kolumny w szablonach formuł
    status: bool = False        # komórka z OK/WARN/ERROR dostaje kolor


def _pct(stawka: Decimal | None) -> Decimal | None:
    return stawka / 100 if stawka is not None else None


def _oplata(kod: str, pole: str) -> Callable[[Zgloszenie, Pozycja], object]:
    def get(_z: Zgloszenie, p: Pozycja) -> object:
        oplata = p.oplata(kod)
        if oplata is None:
            return None
        value = getattr(oplata, pole)
        return _pct(value) if pole == "stawka" else value
    return get


def _kolumny_oplaty(kod: str, nazwa: str) -> list[Kol]:
    return [Kol(f"{kod} podstawa", _oplata(kod, "podstawa"), MONEY, 13, f"{kod}_podstawa"),
            Kol(f"{kod} stawka", _oplata(kod, "stawka"), PCT, 9),
            Kol(f"{kod} kwota wyliczona", _oplata(kod, "kwota_wyliczona"), MONEY, 13),
            Kol(f"{kod} należna ({nazwa})", _oplata(kod, "kwota_nalezna"), MONEY, 13,
                f"{kod}_nalezna"),
            Kol(f"{kod} metoda", _oplata(kod, "metoda"), None, 8)]


def _inne_oplaty(_z: Zgloszenie, p: Pozycja) -> str:
    """Opłaty poza A00/B00 (np. cło antydumpingowe) — tekstem, żeby nie zniknęły z raportu."""
    return "; ".join(f"{o.kod}: {o.kwota_nalezna} PLN" for o in p.oplaty
                     if o.kod not in ("A00", "B00"))


ZGLOSZENIA = [
    Kol("Nr SAD", lambda z, w: z.numer, szer=11),
    Kol("Stan AIS", lambda z, w: z.stan_ais, szer=18),
    Kol("Data zgłoszenia", lambda z, w: z.data_zgloszenia, DATA, 12),
    Kol("Data wydruku", lambda z, w: z.data_wydruku, DATA_CZAS, 16),
    Kol("Nr referencyjny", lambda z, w: z.nr_referencyjny, szer=14),
    Kol("LRN", lambda z, w: z.lrn, szer=14),
    Kol("Nadawca", lambda z, w: z.nadawca, szer=36),
    Kol("Odbiorca", lambda z, w: z.odbiorca, szer=32),
    Kol("Zgłaszający", lambda z, w: z.zglaszajacy, szer=28),
    Kol("Importer NIP", lambda z, w: z.importer_nip, szer=13),
    Kol("Importer EORI", lambda z, w: z.importer_eori, szer=19),
    Kol("Zgłaszający NIP", lambda z, w: z.zglaszajacy_nip, szer=13),
    Kol("Typ zgł. [11 01/02]", lambda z, w: z.typ_zgloszenia, szer=10),
    Kol("UC zgłoszenia", lambda z, w: z.uc_zgloszenia, szer=11),
    Kol("Liczba pozycji", lambda z, w: z.liczba_pozycji, szer=9),
    Kol("Liczba opakowań", lambda z, w: z.liczba_opakowan, szer=10),
    Kol("Kurs [14 09]", lambda z, w: z.kurs, KURS, 9),
    Kol("Eksporter", lambda z, w: z.eksporter, szer=36),
    Kol("Kontenery", lambda z, w: ", ".join(z.kontenery), szer=14),
    Kol("Masa brutto [18 04]", lambda z, w: z.masa_brutto, MASA, 12),
    Kol("Wartość faktur [14 06]", lambda z, w: z.wartosc_faktur, MONEY, 14),
    Kol("Waluta", lambda z, w: z.waluta, szer=7),
    Kol("Kraj przeznaczenia", lambda z, w: z.kraj_przeznaczenia, szer=8),
    Kol("Kraj wysyłki", lambda z, w: z.kraj_wysylki, szer=8),
    Kol("Rodzaj transakcji", lambda z, w: z.rodzaj_transakcji, szer=8),
    Kol("Incoterm", lambda z, w: z.incoterm, szer=8),
    Kol("Miejsce dostawy", lambda z, w: z.miejsce_dostawy, szer=16),
    Kol("Lokalizacja (nr pozwolenia)", lambda z, w: z.lokalizacja_pozwolenie, szer=20),
    Kol("VAT FR7", lambda z, w: z.vat_fr7, szer=14),
    Kol("GRN", lambda z, w: z.grn, szer=19),
    Kol("Σ cło A00 (należne)", lambda z, w: z.suma_naleznych("A00"), MONEY, 13),
    Kol("Σ VAT B00 (należne)", lambda z, w: z.suma_naleznych("B00"), MONEY, 13),
    Kol("Status walidacji", lambda z, w: najgorszy(w), szer=10, status=True),
]

POZYCJE = [
    Kol("Nr SAD", lambda z, p: z.numer, szer=11),
    Kol("LRN", lambda z, p: z.lrn, szer=14),
    Kol("Kontener(y)", lambda z, p: ", ".join(z.kontenery), szer=14),
    Kol("Faktura(y) N935", lambda z, p: ", ".join(p.faktury), szer=16),
    Kol("Proforma(y) N325", lambda z, p: ", ".join(p.proformy), szer=16),
    Kol("Poz.", lambda z, p: p.nr, szer=6),
    Kol("Opis [18 05]", lambda z, p: p.opis, szer=50),
    Kol("CN [18 09]", lambda z, p: p.cn, szer=10),
    Kol("TARIC", lambda z, p: p.taric, szer=7),
    Kol("Kody dodatkowe [18 07]", lambda z, p: ", ".join(p.kody_dodatkowe), szer=16),
    Kol("Dokument poprzedni [12 01]", lambda z, p: p.dokument_poprzedni, szer=30),
    Kol("Wartość fakturowa [14 08]", lambda z, p: p.wartosc_fakturowa, MONEY, 14, "wart_fakt"),
    Kol("Waluta", lambda z, p: z.waluta, szer=7),
    Kol("Kurs [14 09]", lambda z, p: z.kurs, KURS, 9, "kurs"),
    Kol("Metoda wyceny [14 10]", lambda z, p: p.metoda_wyceny, szer=8),
    Kol("Preferencje [14 11]", lambda z, p: p.preferencje, szer=8),
    Kol("Procedura [11 09]", lambda z, p: p.procedura, szer=8),
    Kol("Kraj pochodzenia", lambda z, p: p.kraj_pochodzenia, szer=8),
    Kol("Masa netto [18 01]", lambda z, p: p.masa_netto, MASA, 12),
    Kol("Ilość w jedn. uzup. [18 02]", lambda z, p: p.ilosc_uzup, szer=12),
    Kol("Wartość stat. [99 06]", lambda z, p: p.wartosc_stat, MONEY, 13),
    Kol("ΣAK (doliczenia)", lambda z, p: p.suma_ak, MONEY, 12, "suma_ak"),
    Kol("ΣCA (transport w UE)", lambda z, p: p.suma_ca, MONEY, 12, "suma_ca"),
    Kol("Opakowania liczba", lambda z, p: p.opakowania_liczba, szer=9),
    Kol("Opakowania rodzaj", lambda z, p: p.opakowania_rodzaj, szer=8),
    *_kolumny_oplaty("A00", "cło"),
    *_kolumny_oplaty("B00", "VAT"),
    Kol("Inne opłaty (kod: należna)", _inne_oplaty, szer=16),
    Kol("Kwota ogółem [14 16]", lambda z, p: p.kwota_ogolem, MONEY, 13),
    # kontrola: te same wzory co validator._pozycja, liczone przez Excel z komórek wiersza
    Kol("Wartość celna wyliczona", lambda z, p: _Formula(
        "=ROUND({wart_fakt}*{kurs}+{suma_ak},0)"), MONEY, 14, "celna"),
    Kol("Różnica wart. celnej (A00 − wyliczona)",
        lambda z, p: _Formula("={A00_podstawa}-{celna}"), MONEY, 14),
    Kol("Podstawa VAT wyliczona", lambda z, p: _Formula(
        "=ROUND({A00_podstawa}+{A00_nalezna}+{suma_ca},0)"), MONEY, 14, "vat"),
    Kol("Różnica podst. VAT (B00 − wyliczona)",
        lambda z, p: _Formula("={B00_podstawa}-{vat}"), MONEY, 14),
]

WALIDACJA = [
    Kol("Nr SAD", lambda z, w: z.numer, szer=11),
    Kol("Pozycja", lambda z, w: w.pozycja, szer=8),
    Kol("Reguła", lambda z, w: w.regula, szer=55),
    Kol("Oczekiwane", lambda z, w: w.oczekiwane, szer=24),
    Kol("Odczytane", lambda z, w: w.odczytane, szer=24),
    Kol("Status", lambda z, w: w.status, szer=9, status=True),
]

DOKUMENTY = [
    Kol("Nr SAD", lambda z, p, d: z.numer, szer=11),
    Kol("Poz.", lambda z, p, d: p.nr, szer=6),
    Kol("Kod dokumentu", lambda z, p, d: d.kod, szer=10),
    Kol("Numer", lambda z, p, d: d.numer, szer=34),
]


def _arkusz(wb: Workbook, tytul: str, kolumny: list[Kol], rekordy: Iterable[tuple]) -> None:
    ws = wb.create_sheet(tytul)
    append_row(ws, [k.naglowek for k in kolumny])
    litery = [get_column_letter(i) for i in range(1, len(kolumny) + 1)]
    for litera, kol, cell in zip(litery, kolumny, ws[1], strict=True):
        cell.font = FONT_NAGLOWEK
        ws.column_dimensions[litera].width = kol.szer
    wiersz = 1
    for rekord in rekordy:
        wiersz += 1
        adresy = {k.klucz: f"{dlt}{wiersz}" for k, dlt in zip(kolumny, litery, strict=True) if k.klucz}
        wartosci = [k.wartosc(*rekord) for k in kolumny]
        append_row(ws, [None if v == "" or isinstance(v, _Formula) else v for v in wartosci])
        for nr, (kol, value) in enumerate(zip(kolumny, wartosci, strict=True), 1):
            cell = ws.cell(row=wiersz, column=nr)
            if isinstance(value, _Formula):
                cell.value = value.format(**adresy)
            cell.font = FONT
            if kol.fmt:
                cell.number_format = kol.fmt
            if kol.status and value in FILLS:
                cell.fill = FILLS[str(value)]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{litery[-1]}{wiersz}"


def build_workbook(zgloszenia: list[Zgloszenie], wyniki: list[list[Wynik]]) -> io.BytesIO:
    """xlsx z czterema arkuszami; `wyniki[i]` to walidacja `zgloszenia[i]`."""
    if len(zgloszenia) != len(wyniki):
        raise ValueError("wyniki walidacji muszą odpowiadać zgłoszeniom 1:1")
    wb = Workbook()
    wb.remove(wb.active)
    pary = list(zip(zgloszenia, wyniki, strict=True))
    _arkusz(wb, "Zgłoszenia", ZGLOSZENIA, pary)
    _arkusz(wb, "Pozycje", POZYCJE, ((z, p) for z in zgloszenia for p in z.pozycje))
    _arkusz(wb, "Walidacja", WALIDACJA, ((z, w) for z, lista in pary for w in lista))
    _arkusz(wb, "Dokumenty", DOKUMENTY,
            ((z, p, d) for z in zgloszenia for p in z.pozycje for d in p.dokumenty))
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
