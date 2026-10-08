"""Importy master data: nagłówki zamówień SAP (EKKO), jednostki materiałów (MARM),
stany DLT, porty kontenerowe, dostawcy (LFA1) i wspólny dropzone z auto-rozpoznaniem.

Sub-router bez prefiksu — podpinany w routers/imports.py pod /api/import."""
import io

from fastapi import APIRouter, Depends, HTTPException, Query, Response, UploadFile, status
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import locode
from ..audit import record
from ..config import settings as app_config
from ..database import get_db
from ..deps import Editors as editors
from ..deps import GlobalDataEditors as global_editors
from ..deps import check_global_data_editor, get_company_by_code
from ..importers.lfa1 import apply_plan, build_plan, parse_lfa1, preview
from ..importers.master_data import (
    _parse_cport_rows,
    _parse_dlt_stock_rows,
    _parse_marm_rows,
    detect_master_type,
)
from ..importers.purchasing import _link_container_id, _parse_ekko_rows
from ..invoices import marm_sync
from ..models import Container, ContainerPort, MaterialUnit, SapImport, SapOrder, User, utcnow
from ..models.sap import ACTIVE, MISSING_IN_SAP
from .forwarding import read_upload_capped

router = APIRouter()


# --- wzory plików importu (xlsx do pobrania przy każdym imporcie) ---

@router.get("/templates/{kind}")
def import_template(kind: str, user: User = editors):
    from ..exports import XLSX_MIME
    from ..import_templates import TEMPLATES, build_template
    if kind not in TEMPLATES:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Nieznany wzór importu")
    return Response(build_template(kind), media_type=XLSX_MIME,
                    headers={"Content-Disposition": f'attachment; filename="wzor_{kind}.xlsx"'})


def _mark_missing(existing: dict, seen: set, full_export: bool) -> int:
    """DATA-003: rekordy spoza PEŁNEGO eksportu → „brak_w_sap” (bez kasowania);
    przy częściowym pliku nic nie zmieniamy. Zwraca liczbę nowo oznaczonych."""
    if not full_export:
        return 0
    missing = [row for key, row in existing.items()
               if key not in seen and row.sap_status != MISSING_IN_SAP]
    for row in missing:
        row.sap_status = MISSING_IN_SAP
    return len(missing)


# --- import nagłówków zamówień (EKKO) z eksportu SAP ---

@router.post("/sap-orders")
def import_sap_orders(
    file: UploadFile,
    company_code: str = Query(...),
    dry_run: bool = Query(default=True),
    full_export: bool = Query(default=False),
    db: Session = Depends(get_db),
    user: User = editors,
):
    """Import nagłówków zamówień zakupu (EKKO) z SAP — „co płynie, skąd i kiedy”.

    Upsert po (spółka, order_number); eksport jest dogrywany sukcesywnie, więc ponowny
    import aktualizuje istniejące zamówienia. Auto-link do kontenera po numerze zamówienia
    (ta sama reguła całych tokenów co przy imporcie ETD). `full_export` — plik to pełny
    eksport EKKO: zamówienia spółki spoza pliku → `sap_status = brak_w_sap` (DATA-003)."""
    company = get_company_by_code(db, user, company_code)
    parsed = _parse_ekko_rows(
        read_upload_capped(file, app_config.max_upload_mb, "Plik EKKO"))

    existing = {o.order_number: o for o in db.scalars(
        select(SapOrder).where(SapOrder.company_id == company.id))}
    containers = [(cid, no) for cid, no in db.execute(select(
        Container.id, Container.order_numbers).where(Container.company_id == company.id))]

    rows: list[tuple[dict, int | None]] = []
    seen: set[str] = set()
    dup_n = 0
    for raw in parsed:
        if raw["order_number"] in seen:
            dup_n += 1
            continue
        seen.add(raw["order_number"])
        rows.append((raw, _link_container_id(raw["order_number"], containers)))

    new_n = sum(1 for raw, _ in rows if raw["order_number"] not in existing)
    # DATA-003: spoza pliku → „brak_w_sap” tylko przy pełnym eksporcie; nic nie kasujemy
    counts = {"total": len(parsed), "new": new_n, "updated": len(rows) - new_n,
              "linked": sum(1 for _, cid in rows if cid), "duplicate": dup_n,
              "disappeared": sum(1 for no in existing if no not in seen)}
    if dry_run:
        return {"dry_run": True, "counts": counts}

    note = f"import EKKO ({file.filename})"
    for raw, container_id in rows:
        order = existing.get(raw["order_number"])
        if order is None:
            order = SapOrder(company_id=company.id, order_number=raw["order_number"])
            db.add(order)
            existing[raw["order_number"]] = order
        for key, value in raw.items():
            if key != "order_number":
                setattr(order, key, value)
        order.sap_status = ACTIVE
        # brak dopasowania nie zrywa istniejącego dowiązania; zmiana dowiązania → audyt
        if container_id is not None and container_id != order.container_id:
            if order.id is not None:
                record(db, entity_type="sap_orders", entity_id=order.id, field="container_id",
                       old_value=order.container_id, new_value=container_id,
                       user=user, note=note)
            order.container_id = container_id
    counts["deactivated"] = _mark_missing(existing, seen, full_export)
    record(db, entity_type="sap_orders", entity_id=company.id, field="import",
           old_value=None,
           new_value=f"{counts['new']} nowych / {counts['updated']} zaktualizowanych",
           user=user, note=note)
    db.add(SapImport(kind="ekko", filename=(file.filename or "")[:255], user_id=user.id,
                     counts=counts, errors=[]))
    db.commit()
    return {"dry_run": False, "counts": counts}


# --- import jednostek materiałów (MARM) z eksportu SAP ---

@router.post("/material-units")
def import_material_units(
    file: UploadFile,
    dry_run: bool = Query(default=True),
    full_export: bool = Query(default=False),
    db: Session = Depends(get_db),
    user: User = global_editors,
):
    """Import przeliczników jednostek materiałów (MARM) z SAP — master data globalne.

    Upsert po (material_no, unit); ponowny import aktualizuje istniejące wiersze.
    Placeholdery (jednostka „0”) są pomijane i policzone. `full_export` — plik to pełny
    eksport MARM: jednostki spoza pliku → `sap_status = brak_w_sap` (DATA-003)."""
    errors: list[dict] = []
    parsed, placeholders = _parse_marm_rows(
        read_upload_capped(file, app_config.max_upload_mb, "Plik MARM"), errors)

    existing = {(m.material_no, m.unit): m for m in db.scalars(select(MaterialUnit))}
    seen: set[tuple[str, str]] = set()
    rows: list[dict] = []
    dup_n = 0
    for raw in parsed:
        key = (raw["material_no"], raw["unit"])
        if key in seen:
            dup_n += 1
            continue
        seen.add(key)
        rows.append(raw)
    new_n = sum(1 for r in rows if (r["material_no"], r["unit"]) not in existing)
    counts = {"total": len(parsed) + placeholders + len(errors), "new": new_n,
              "updated": len(rows) - new_n, "placeholders": placeholders,
              "duplicate": dup_n, "invalid": len(errors),
              "disappeared": sum(1 for key in existing if key not in seen)}
    # materiały dla dopasowania faktur (numer, jm podstawowa, EAN, przeliczniki)
    counts.update(marm_sync.sync(db, rows, dry_run))
    if dry_run:
        return {"dry_run": True, "counts": counts, "errors": errors[:200]}

    for raw in rows:
        unit_row = existing.get((raw["material_no"], raw["unit"]))
        if unit_row is None:
            unit_row = MaterialUnit(material_no=raw["material_no"], unit=raw["unit"])
            db.add(unit_row)
            existing[(raw["material_no"], raw["unit"])] = unit_row
        for key, value in raw.items():
            if key not in ("material_no", "unit"):
                setattr(unit_row, key, value)
        unit_row.sap_status = ACTIVE
    counts["deactivated"] = _mark_missing(existing, seen, full_export)
    record(db, entity_type="material_units", entity_id=0, field="import",
           old_value=None,
           new_value=f"{counts['new']} nowych / {counts['updated']} zaktualizowanych",
           user=user, note=f"import MARM ({file.filename})")
    db.add(SapImport(kind="marm", filename=(file.filename or "")[:255], user_id=user.id,
                     counts=counts, errors=errors))
    db.commit()
    return {"dry_run": False, "counts": counts, "errors": errors[:200]}


# --- import stanów DLT z xlsx (fallback gdy Power BI ich nie podaje) ---

@router.post("/dlt-stock")
def import_dlt_stock(
    file: UploadFile,
    dry_run: bool = Query(default=True),
    db: Session = Depends(get_db),
    user: User = global_editors,
):
    """Import stanów DLT z xlsx — snapshot (każdy import kasuje poprzedni).

    Analityka wywołań używa tych stanów, gdy Power BI ich nie podaje (fallback)."""
    parsed = _parse_dlt_stock_rows(
        read_upload_capped(file, app_config.max_upload_mb, "Plik stanów DLT"))
    counts = {"rows": len(parsed),
              "with_hu": sum(1 for r in parsed if r["hu"]),
              "produkty": len({r["produkt"] for r in parsed})}
    if dry_run:
        return {"dry_run": True, "counts": counts}

    from ..models import DltStock
    db.query(DltStock).delete()   # snapshot: zastępujemy poprzedni import w całości
    now = utcnow()
    for raw in parsed:
        db.add(DltStock(produkt=raw["produkt"], ilosc=raw["ilosc"], hu=raw["hu"],
                        lokalizacja=raw["lokalizacja"], source_file=file.filename or "",
                        imported_at=now))
    record(db, entity_type="dlt_stock", entity_id=0, field="import", old_value=None,
           new_value=f"{counts['rows']} wierszy / {counts['produkty']} produktów",
           user=user, note=f"import stanów DLT ({file.filename})")
    db.commit()
    # świeży snapshot → przelicz analizę przy następnym odczycie
    from ..pallets_cache import invalidate
    invalidate(db)
    return {"dry_run": False, "counts": counts}


# --- import portów kontenerowych (słownik globalny, warstwa mapy trackingu) ---

@router.post("/container-ports")
def import_container_ports(
    file: UploadFile,
    dry_run: bool = Query(default=True),
    db: Session = Depends(get_db),
    user: User = global_editors,
):
    """Import słownika portów kontenerowych (xlsx albo tsv/csv). Master data globalne.

    Idempotentny upsert po code; współrzędne dopasowywane z zasobu UN/LOCODE
    (kod wprost → kraj+nazwa → fuzzy w kraju); brak dopasowania → lat/lon NULL."""
    parsed = _parse_cport_rows(
        read_upload_capped(file, app_config.max_upload_mb, "Plik portów"),
        file.filename or "")

    seen: set[str] = set()
    rows: list[dict] = []
    dup_n = 0
    for raw in parsed:
        if raw["code"] in seen:
            dup_n += 1
            continue
        seen.add(raw["code"])
        coords = locode.match_coords(raw["code"], raw["name"], raw["country_code"])
        raw["lat"], raw["lon"] = coords if coords else (None, None)
        rows.append(raw)

    existing = {p.code: p for p in db.scalars(select(ContainerPort))}
    new_n = sum(1 for r in rows if r["code"] not in existing)
    with_coords = sum(1 for r in rows if r["lat"] is not None)
    counts = {"total": len(parsed), "new": new_n, "updated": len(rows) - new_n,
              "with_coords": with_coords, "without_coords": len(rows) - with_coords,
              "duplicate": dup_n}
    if dry_run:
        return {"dry_run": True, "counts": counts}

    for raw in rows:
        port = existing.get(raw["code"])
        if port is None:
            port = ContainerPort(code=raw["code"])
            db.add(port)
            existing[raw["code"]] = port
        for key, value in raw.items():
            if key != "code":
                setattr(port, key, value)
    record(db, entity_type="container_ports", entity_id=0, field="import",
           old_value=None,
           new_value=f"{counts['new']} nowych / {counts['updated']} zaktualizowanych"
                     f" / {with_coords} ze współrzędnymi",
           user=user, note=f"import portów kontenerowych ({file.filename})")
    db.commit()
    return {"dry_run": False, "counts": counts}


# --- import dostawców (LFA1) z eksportu SAP ---

def _require_catalog(user: User) -> None:
    from ..deps import supplier_catalog_access
    if not supplier_catalog_access(user):
        raise HTTPException(status.HTTP_403_FORBIDDEN,
                            "Import kartoteki dostawców — tylko konta z dostępem do kartoteki.")


@router.post("/suppliers-lfa1")
def import_suppliers_lfa1(
    file: UploadFile,
    dry_run: bool = Query(default=True),
    full_export: bool = Query(default=False),
    db: Session = Depends(get_db),
    user: User = editors,
):
    """Import kartoteki dostawców z eksportu SAP LFA1 z podglądem (importers/lfa1.py).
    `full_export` — plik to pełny eksport LFA1: brakujący w nim dostawcy → nieaktywni w SAP."""
    _require_catalog(user)
    rows, errors = parse_lfa1(read_upload_capped(file, app_config.max_upload_mb, "Plik LFA1"))
    plan = build_plan(db, rows, errors)
    result = preview(plan)     # przed zapisem — „stara” wartość w zmianach
    if dry_run:
        return {"dry_run": True, **result}
    log = apply_plan(db, plan, full_export=full_export, user=user, filename=file.filename or "")
    db.commit()
    return {"dry_run": False, "import_id": log.id, "deactivated": log.counts["deactivated"],
            **result}


@router.get("/sap-imports/{import_id}/errors.xlsx")
def sap_import_errors(import_id: int, db: Session = Depends(get_db), user: User = editors):
    """Raport błędów importu SAP (wiersze pominięte) jako xlsx."""
    _require_catalog(user)
    log = db.get(SapImport, import_id)
    if log is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Nie ma takiego importu.")
    from openpyxl import Workbook

    from ..exports import append_row
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Błędy"
    append_row(sheet, ["Wiersz", "Kod SAP", "Nazwa", "Powód"])
    for err in log.errors or []:   # wartości z wgranego pliku — tylko jako tekst (SEC-008)
        append_row(sheet, [err.get("row"), err.get("sap_code"), err.get("name"),
                           err.get("reason")])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return Response(buffer.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition":
                             f'attachment; filename="import-{log.kind}-{log.id}-bledy.xlsx"'})


# --- auto-rozpoznanie pliku master data (#18): jeden dropzone, typ po nagłówkach ---

@router.post("/master-data")
def import_master_data(
    file: UploadFile,
    company_code: str | None = Query(default=None),
    dry_run: bool = Query(default=True),
    full_export: bool = Query(default=False),   # LFA1 / EKKO / MARM — pełny eksport SAP
    db: Session = Depends(get_db),
    user: User = editors,
):
    """Wspólny dropzone „Wgraj plik master data": typ rozpoznany po nagłówkach,
    plik kierowany do właściwego importera. Niepewność → 422 z listą kandydatów."""
    content = read_upload_capped(file, app_config.max_upload_mb, "Plik master data")
    detected, candidates = detect_master_type(content, file.filename or "")
    if detected is None:
        hint = f" Częściowo pasuje do: {', '.join(candidates)}." if candidates else ""
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            "Nie rozpoznano typu pliku master data po nagłówkach"
                            f" (MARM / EKKO / LFA1 / porty kontenerowe).{hint}")
    upload = UploadFile(io.BytesIO(content), filename=file.filename)
    if detected in ("marm", "cports"):
        check_global_data_editor(user)   # dane wspólne grupy — jak trasy bezpośrednie (ACL-003)
    if detected == "marm":
        result = import_material_units(upload, dry_run, full_export, db, user)
    elif detected == "cports":
        result = import_container_ports(upload, dry_run, db, user)
    elif detected == "lfa1":
        result = import_suppliers_lfa1(upload, dry_run, full_export, db, user)
    else:  # ekko wymaga spółki (zamówienia są per spółka)
        if not company_code:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                                "Rozpoznano eksport EKKO — wybierz spółkę (company_code).")
        result = import_sap_orders(upload, company_code, dry_run, full_export, db, user)
    return {"detected": detected, **result}
