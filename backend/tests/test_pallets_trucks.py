"""Wywołania-DLT: auta 4×33, MARM, cele dni per materiał, mail, Excel per auto,
degradacja feature-detect, izolacja spółek."""
import io

import openpyxl
import pytest

from app.pallets_analysis import build_rows


@pytest.fixture(autouse=True)
def _mock_powerbi(monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "powerbi_provider", "mock")
    monkeypatch.setattr(settings, "powerbi_dlt_locations", "3DLT")
    monkeypatch.setattr(settings, "powerbi_table_stock", "stock")
    monkeypatch.setattr(settings, "powerbi_table_vbba", "vbba")
    monkeypatch.setattr(settings, "powerbi_table_orders", "vbbe")
    monkeypatch.setattr(settings, "powerbi_table_usage", "zuzycie")
    monkeypatch.setattr(settings, "powerbi_cache_ttl_min", 0)


def _company_id(client, headers):
    return client.get("/api/companies", headers=headers).json()[0]["id"]


def _create(client, headers, lines, **extra):
    body = {"company_id": _company_id(client, headers), "allow_over_dlt": True,
            "lines": lines, **extra}
    return client.post("/api/pallet-calls", headers=headers, json=body)


# --- guardy 33 / 132 ---

def test_guard_33_per_truck(client, admin_headers):
    r = _create(client, admin_headers,
                [{"produkt": "A", "ilosc_pal": 34, "truck_no": 1}])
    assert r.status_code == 422 and "33" in r.text
    r2 = _create(client, admin_headers,
                 [{"produkt": "A", "ilosc_pal": 20, "truck_no": 1},
                  {"produkt": "B", "ilosc_pal": 14, "truck_no": 1}])
    assert r2.status_code == 422
    r3 = _create(client, admin_headers,
                 [{"produkt": "A", "ilosc_pal": 20, "truck_no": 1},
                  {"produkt": "B", "ilosc_pal": 13, "truck_no": 1}])
    assert r3.status_code == 201, r3.text
    assert r3.json()["trucks"] == [
        {"id": r3.json()["trucks"][0]["id"], "ordinal": 1, "capacity": 33}]


def test_guard_132_total(client, admin_headers):
    r = _create(client, admin_headers, [{"produkt": "A", "ilosc_pal": 133}])
    assert r.status_code == 422 and "132" in r.text


def test_guard_ceil_of_fractional_pallets(client, admin_headers):
    # 32.2 palety ułamkowo -> 33 miejsca -> OK; 33.1 -> 34 -> 422
    ok = _create(client, admin_headers,
                 [{"produkt": "A", "ilosc_pal": 33, "pallets": 32.2, "truck_no": 2}])
    assert ok.status_code == 201, ok.text
    bad = _create(client, admin_headers,
                  [{"produkt": "A", "ilosc_pal": 33, "pallets": 33.1, "truck_no": 2}])
    assert bad.status_code == 422


def test_truck_no_out_of_range_rejected(client, admin_headers):
    r = _create(client, admin_headers,
                [{"produkt": "A", "ilosc_pal": 1, "truck_no": 5}])
    assert r.status_code == 422


# --- MARM: palety ułamkowe i brak PAL ---

def test_marm_pal_conversion_and_paz_override(client, admin_headers, db_session):
    from app.models import MaterialUnit
    db_session.add(MaterialUnit(material_no="DEMO-SKU-079", unit="PAL",
                                numerator=10, denominator=1))
    db_session.commit()
    r = client.get("/api/pallet-calls/analysis?q=DEMO-SKU-079&refresh=true",
                   headers=admin_headers)
    item = next(i for i in r.json()["items"] if i["produkt"] == "DEMO-SKU-079")
    assert item["stan_mag_pal"] == pytest.approx(7.9)     # 79 szt / 10 szt/pal
    assert item["sugestia_pal"] is not None
    # słownik PAZ nadpisuje MARM
    client.post("/api/paz", headers=admin_headers,
                json={"produkt": "DEMO-SKU-079", "sztuk_na_palete": 79})
    r2 = client.get("/api/pallet-calls/analysis?q=DEMO-SKU-079&refresh=true",
                    headers=admin_headers)
    item2 = next(i for i in r2.json()["items"] if i["produkt"] == "DEMO-SKU-079")
    assert item2["stan_mag_pal"] == pytest.approx(1.0)


def test_no_pal_unit_keeps_szt_only(client, admin_headers):
    r = client.get("/api/pallet-calls/analysis?q=DEMO-SKU-079&refresh=true",
                   headers=admin_headers)
    item = next(i for i in r.json()["items"] if i["produkt"] == "DEMO-SKU-079")
    assert item["sugestia_pal"] is None and item["stan_mag_pal"] is None
    assert item["stan_mag_szt"] == 79      # sztuki zawsze widoczne
    assert any(d["powod"] == "brak_paz" and d["produkt"] == "DEMO-SKU-079"
               for d in r.json()["discrepancies"])


# --- degradacja feature-detect (bez kolumn HU / rozchodów dziennych) ---

_KW = dict(key="produkt", location_field="lok", dlt_values={"DLT"},
           open_statuses=set(), target_days=14, urgent_threshold=2)


def test_degradation_without_hu_column():
    stock = [{"produkt": "A", "ilosc": 100, "lok": "DLT"}]
    rows, _, feats = build_rows(stock, [], [], [], {"A": 10}, {}, **_KW)
    assert feats["hu"] is False and rows[0]["hu"] == []


def test_hu_column_detected_and_grouped():
    stock = [{"produkt": "A", "ilosc": 60, "lok": "DLT", "glowna_hu": "H1"},
             {"produkt": "A", "ilosc": 40, "lok": "DLT", "glowna_hu": "H2"},
             {"produkt": "A", "ilosc": 5, "lok": "MAG", "glowna_hu": "H3"}]
    rows, _, feats = build_rows(stock, [], [], [], {"A": 10}, {}, **_KW)
    assert feats["hu"] is True
    assert rows[0]["hu"] == [{"hu": "H1", "ilosc": 60}, {"hu": "H2", "ilosc": 40}]


def test_daily_usage_rolling_average_vs_legacy():
    stock = [{"produkt": "A", "ilosc": 10, "lok": "MAG"}]
    legacy = [{"produkt": "A", "zuzycie_dzienne": 5}]
    rows, _, feats = build_rows(stock, [], [], legacy, {"A": 1}, {}, **_KW)
    assert feats["daily_usage"] is False
    assert rows[0]["zuzycie_szt_dzien"] == 5
    daily = [{"produkt": "A", "data": "2026-09-01", "ilosc": 60},
             {"produkt": "A", "data": "2026-09-02", "ilosc": 30}]
    rows2, _, feats2 = build_rows(stock, [], [], daily, {"A": 1}, {}, **_KW)
    assert feats2["daily_usage"] is True
    assert rows2[0]["zuzycie_szt_dzien"] == pytest.approx(90 / 30)


def test_analysis_endpoint_exposes_features(client, admin_headers):
    r = client.get("/api/pallet-calls/analysis?refresh=true", headers=admin_headers)
    assert r.json()["features"] == {"hu": True, "daily_usage": False}


# --- cel dni: globalny + nadpisanie per materiał ---

def test_stock_target_overrides_global(client, admin_headers):
    client.post("/api/paz", headers=admin_headers,
                json={"produkt": "DEMO-SKU-020", "sztuk_na_palete": 100})
    base = client.get("/api/pallet-calls/analysis?q=DEMO-SKU-020&refresh=true",
                      headers=admin_headers).json()["items"][0]
    assert base["cel_dni"] == 14 and base["sugestia_pal"] == 10
    r = client.post("/api/stock-targets", headers=admin_headers,
                    json={"material_no": "DEMO-SKU-020", "days": 28})
    assert r.status_code == 200
    after = client.get("/api/pallet-calls/analysis?q=DEMO-SKU-020&refresh=true",
                       headers=admin_headers).json()["items"][0]
    assert after["cel_dni"] == 28
    # potrzeba 28*1 + 6 - 10 = 24, przycięte do stanu DLT (20 pal)
    assert after["sugestia_pal"] == 20
    lst = client.get("/api/stock-targets", headers=admin_headers).json()
    assert lst["global_days"] == 14
    target_id = lst["items"][0]["id"]
    assert client.delete(f"/api/stock-targets/{target_id}",
                         headers=admin_headers).status_code == 204


def test_stock_target_upsert_updates_days(client, admin_headers):
    client.post("/api/stock-targets", headers=admin_headers,
                json={"material_no": "M-1", "days": 7})
    client.post("/api/stock-targets", headers=admin_headers,
                json={"material_no": "M-1", "days": 21})
    items = client.get("/api/stock-targets", headers=admin_headers).json()["items"]
    assert [i for i in items if i["material_no"] == "M-1"][0]["days"] == 21
    assert len([i for i in items if i["material_no"] == "M-1"]) == 1


# --- mail przy wysłaniu ---

def test_send_uses_dlt_warehouse_email_and_cc_author(client, admin_headers,
                                                     monkeypatch, db_session):
    from app.models import User, Warehouse
    from app.routers import pallets
    cid = _company_id(client, admin_headers)
    db_session.add(Warehouse(name="Magazyn DLT", email="dlt@dlt.pl", company_id=cid))
    admin = db_session.query(User).filter_by(login="admin").first()
    admin.email = "autor@example.com"
    db_session.commit()
    sent = {}
    monkeypatch.setattr(pallets, "send_html_email",
                        lambda recipients, *a, **k: sent.update(
                            {"recipients": recipients, "kw": k}))
    call = _create(client, admin_headers,
                   [{"produkt": "X", "ilosc_pal": 1, "truck_no": 1}]).json()
    r = client.post(f"/api/pallet-calls/{call['id']}/send", headers=admin_headers)
    assert r.status_code == 200 and r.json()["status"] == "sent"
    assert r.json()["sent_at"] is not None
    assert sent["recipients"] == ["dlt@dlt.pl", "autor@example.com"]
    assert sent["kw"]["attachments"][0][0].endswith(".xlsx")


def test_send_fallback_env_emails(client, admin_headers, monkeypatch):
    from app.config import settings
    from app.routers import pallets
    monkeypatch.setattr(settings, "dlt_call_emails", "a@dlt.pl, b@dlt.pl")
    sent = {}
    monkeypatch.setattr(pallets, "send_html_email",
                        lambda recipients, *a, **k: sent.update({"r": recipients}))
    call = _create(client, admin_headers, [{"produkt": "X", "ilosc_pal": 1}]).json()
    r = client.post(f"/api/pallet-calls/{call['id']}/send", headers=admin_headers)
    assert r.status_code == 200
    assert sent["r"][:2] == ["a@dlt.pl", "b@dlt.pl"]


def test_send_without_any_address_400(client, admin_headers):
    call = _create(client, admin_headers, [{"produkt": "X", "ilosc_pal": 1}]).json()
    r = client.post(f"/api/pallet-calls/{call['id']}/send", headers=admin_headers)
    assert r.status_code == 400


# --- Excel: arkusz per auto ---

def test_xlsx_sheet_per_truck(client, admin_headers):
    call = _create(client, admin_headers, [
        {"produkt": "A", "ilosc_pal": 3, "pallets": 2.5, "truck_no": 1,
         "hu_numbers": "H1,H2,H3"},
        {"produkt": "B", "ilosc_pal": 4, "truck_no": 2},
        {"produkt": "C", "ilosc_pal": 1},
    ]).json()
    x = client.get(f"/api/pallet-calls/{call['id']}/xlsx", headers=admin_headers)
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Auto 1", "Auto 2", "Bez auta"]
    ws = wb["Auto 1"]
    row = list(ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]
    assert row[0] == 1 and row[1] == "A" and row[3] == "H1,H2,H3"
    assert row[4] == 2.5 and row[5] == 3


def test_xlsx_legacy_single_sheet_without_trucks(client, admin_headers):
    call = _create(client, admin_headers, [{"produkt": "A", "ilosc_pal": 2}]).json()
    x = client.get(f"/api/pallet-calls/{call['id']}/xlsx", headers=admin_headers)
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Wywolanie"]


# --- izolacja spółek ---

def test_company_isolation(client, admin_headers):
    companies = client.get("/api/companies", headers=admin_headers).json()
    other = companies[1]
    client.post("/api/users", headers=admin_headers, json={
        "login": "log.other", "password": "haslo123", "role": "logistics",
        "company_id": other["id"], "view_all_companies": False})
    call = _create(client, admin_headers, [{"produkt": "X", "ilosc_pal": 1}]).json()
    login = client.post("/api/auth/login",
                        data={"username": "log.other", "password": "haslo123"})
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    r = client.get(f"/api/pallet-calls/{call['id']}", headers=headers)
    assert r.status_code in (403, 404)
    lst = client.get("/api/pallet-calls", headers=headers).json()
    assert all(c["id"] != call["id"] for c in lst)
