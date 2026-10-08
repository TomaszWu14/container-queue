import io

import openpyxl

from app import models
from app.pallets_export import build_xlsx


def test_build_xlsx_has_lines():
    call = models.PalletCall(number="PC-2026-0001", status=models.PalletCallStatus.draft)
    call.lines = [
        models.PalletCallLine(produkt="DEMO-SKU-020", krotki_opis="Strzykawka",
                              ilosc_pal=3),
    ]
    data = build_xlsx(call)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    header = [str(c.value).lower() for c in ws[2]]
    assert "produkt" in header and "ilość palet" in header
    values = [c.value for c in ws[3]]
    assert "DEMO-SKU-020" in values and 3 in values
