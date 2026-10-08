"""Dane materiałowe SharePoint (2026-09-29): jeden plik → zakładka = tabela; techniczne i ukryte
zakładki pomijane; „Hierarchia produktów” daje materiałom KRÓTKĄ nazwę PL, EN i CN."""
import io

from openpyxl import Workbook
from sqlalchemy import select

from app.database import SessionLocal
from app.models import Material


def _file() -> bytes:
    wb = Workbook()
    wb.active.title = "Admin"                                 # techniczna — bez nagłówka
    hier = wb.create_sheet("Hierarchia produktów")
    hier.append(["H1", "REF", "TXT_LONG_PL", "TXT_SHORT_PL", "TXT_SHORT_EN", "KOD_CN", None])
    hier.append(["02", "AT-SGS-XL_1", "Fartuch chirurgiczny STANDARD", "Fart.STANDARD_XL",
                 "Surg. gown XL", "62101092", None])
    hier.append(["02", "EXISTING-1", "Długa", "Krótka nowa", "", "", None])
    blz = wb.create_sheet("BLOZ")
    blz.append(["MATNR", "REF", "JM", "EAN"])
    blz.append(["000000000001108962", "ZCP-02-2012", "SZT", "5900000000001"])
    hidden = wb.create_sheet("Dane_Xpertis")
    hidden.append(["MATNR", "REF"])
    hidden.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client, headers, dry_run):
    return client.post(f"/api/import/sp-materials?dry_run={str(dry_run).lower()}", headers=headers,
                       files={"file": ("SAP_Dane_materiałowe.xlsm", io.BytesIO(_file()),
                                       "application/vnd.ms-excel.sheet.macroEnabled.12")})


def test_import_sheets_and_materials(client, admin_headers):
    with SessionLocal() as db:
        db.add(Material(ref_code="EXISTING-1", ref_norm="EXISTING1", name_pl="stara", tariff_cn="1234"))
        db.commit()
    dry = _post(client, admin_headers, True).json()
    assert dry["counts"]["sheets"] == [{"name": "Hierarchia produktów", "rows": 2},
                                       {"name": "BLOZ", "rows": 1}]
    assert (dry["counts"]["materials_new"], dry["counts"]["materials_updated"]) == (1, 1)
    assert client.get("/api/sp-materials/sheets", headers=admin_headers).json() == []

    assert _post(client, admin_headers, False).status_code == 200
    with SessionLocal() as db:
        gown = db.scalar(select(Material).where(Material.ref_code == "AT-SGS-XL_1"))
        assert (gown.name_pl, gown.name_en, gown.tariff_cn, gown.ref_norm) == \
            ("Fart.STANDARD_XL", "Surg. gown XL", "62101092", "ATSGSXL1")
        old = db.scalar(select(Material).where(Material.ref_code == "EXISTING-1"))
        assert (old.name_pl, old.tariff_cn) == ("Krótka nowa", "1234")   # pusty CN nie czyści

    sheets = client.get("/api/sp-materials/sheets", headers=admin_headers).json()
    assert [(s["name"], s["rows"]) for s in sheets] == [("Hierarchia produktów", 2), ("BLOZ", 1)]
    table = client.get("/api/sp-materials/sheets/Hierarchia produktów?q=fart", headers=admin_headers).json()
    assert table["headers"][:2] == ["H1", "REF"] and len(table["headers"]) == 6
    assert table["total"] == 1 and table["rows"][0][1] == "AT-SGS-XL_1"
    # ponowny import podmienia arkusz, nie dubluje wierszy
    _post(client, admin_headers, False)
    assert client.get("/api/sp-materials/sheets/BLOZ", headers=admin_headers).json()["total"] == 1


def test_sheet_404(client, admin_headers):
    assert client.get("/api/sp-materials/sheets/Nie ma", headers=admin_headers).status_code == 404


def test_import_in_background_thread_and_status(client, admin_headers, monkeypatch):
    """2026-09-29: pełny plik na produkcji przekraczał 30 s żądania — import w wątku tła,
    front odpytuje /import-status (running → done)."""
    import threading

    from app.config import settings
    monkeypatch.setattr(settings, "run_background_jobs", True)
    started = []
    real_thread = threading.Thread

    def capture(*args, **kwargs):
        thread = real_thread(*args, **kwargs)
        started.append(thread)
        return thread
    monkeypatch.setattr(threading, "Thread", capture)
    assert client.get("/api/sp-materials/import-status", headers=admin_headers).json() == {"status": "none"}
    first = _post(client, admin_headers, False).json()
    assert first["status"] in ("running", "done") and started
    started[0].join(timeout=30)
    status = client.get("/api/sp-materials/import-status", headers=admin_headers).json()
    assert status["status"] == "done" and status["counts"]["materials_new"] == 2
    assert status["filename"] == "SAP_Dane_materiałowe.xlsm"


def test_import_error_is_reported_in_status(client, admin_headers):
    bad = {"file": ("zly.xlsm", io.BytesIO(b"to nie jest excel"), "application/octet-stream")}
    r = client.post("/api/import/sp-materials?dry_run=false", headers=admin_headers, files=bad)
    assert r.status_code == 200 and r.json()["status"] == "error"
    assert "odczytać pliku" in r.json()["error"]
    status = client.get("/api/sp-materials/import-status", headers=admin_headers).json()
    assert status["status"] == "error"
