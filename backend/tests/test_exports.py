"""Strażnik: każdy eksport xlsx/csv backendu neutralizuje wstrzyknięcie formuł (app/exports.py)."""
import io

import openpyxl

from app import models
from app.exports import append_row, csv_safe
from app.pallets_export import build_xlsx
from tests.test_notifications import _setup

EVIL = '=HYPERLINK("http://x","klik")'


def _reload(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data))


def test_append_row_keeps_formula_text_as_text():
    wb = openpyxl.Workbook()
    ws = wb.active
    append_row(ws, [EVIL, 5, 1.5, None])
    buf = io.BytesIO()
    wb.save(buf)
    cells = _reload(buf.getvalue()).active[1]
    assert cells[0].value == EVIL and cells[0].data_type == "s"
    assert cells[1].value == 5 and cells[2].value == 1.5


def test_csv_safe():
    assert csv_safe("=1+1") == "'=1+1"
    assert csv_safe("@x") == "'@x" and csv_safe("+x") == "'+x" and csv_safe("-x") == "'-x"
    assert csv_safe("abc") == "abc" and csv_safe(None) == "" and csv_safe(7) == "7"


def test_pallet_call_xlsx_note_is_not_formula():
    call = models.PalletCall(number="PC-2026-0001", status=models.PalletCallStatus.draft)
    call.lines = [models.PalletCallLine(produkt="P1", krotki_opis="x", ilosc_pal=1, note=EVIL)]
    ws = _reload(build_xlsx(call)).active
    cell = ws[3][4]
    assert cell.value == EVIL and cell.data_type == "s"


def test_pallet_call_xlsx_with_trucks_note_is_not_formula():
    call = models.PalletCall(number="PC-2026-0002", status=models.PalletCallStatus.draft)
    truck = models.PalletCallTruck(id=1, ordinal=1)
    call.trucks = [truck]
    call.lines = [models.PalletCallLine(produkt="P1", krotki_opis="x", ilosc_pal=1,
                                        truck_id=1, note=EVIL)]
    ws = _reload(build_xlsx(call)).active
    cell = ws[3][7]
    assert cell.value == EVIL and cell.data_type == "s"


def test_queue_export_notes_are_not_formula(client, admin_headers):
    ctx = _setup(client, admin_headers)
    cid = ctx["container"]["id"]
    assert client.patch(f"/api/containers/{cid}", headers=admin_headers,
                        json={"notes": EVIL}).status_code == 200
    response = client.get("/api/containers/export/xlsx", headers=admin_headers)
    assert response.status_code == 200
    ws = _reload(response.content).active
    notes = [row[-1] for row in ws.iter_rows(min_row=2)]
    assert any(c.value == EVIL and c.data_type == "s" for c in notes)
    assert not any(c.data_type == "f" for c in notes)


def test_append_row_is_linear_not_quadratic():
    # audyt PERF-002: stara wersja szukała dołożonego wiersza przez ws[ws.max_row] (skan całego
    # arkusza przy każdym wierszu) — 4000 wierszy × 30 kolumn ≈ 19 s; liniowo ≈ 0,2 s
    import time
    wb = openpyxl.Workbook()
    ws = wb.active
    start = time.perf_counter()
    for i in range(4000):
        append_row(ws, [f"MSCU{i:07d}", i, EVIL] + [f"t{j}" for j in range(27)])
    # próg z zapasem na CI równoległe (-n auto + pokrycie): liniowo ~0,2 s, kwadratowo ~19 s
    assert time.perf_counter() - start < 10
    assert ws.max_row == 4000 and ws.cell(4000, 3).data_type == "s" and ws.cell(4000, 2).value == 3999


def test_append_row_keeps_dates_as_dates():
    import datetime
    wb = openpyxl.Workbook()
    ws = wb.active
    append_row(ws, [datetime.date(2026, 9, 28), datetime.datetime(2026, 9, 28, 12, 30)])
    buf = io.BytesIO()
    wb.save(buf)
    cells = _reload(buf.getvalue()).active[1]
    assert cells[0].value.date() == datetime.date(2026, 9, 28) and cells[1].value.hour == 12


def test_csv_safe_leading_whitespace_and_control_chars():
    # audyt SEC-008: Excel pomija wiodące białe znaki, a \t i \r OWASP liczy jako wyzwalacze
    for evil in ("\t=1+1", "\r=1+1", " =1+1", "\n@x", "  -2+3"):
        assert csv_safe(evil) == "'" + evil
    assert csv_safe("\tabc") == "'\tabc" and csv_safe(" abc") == " abc" and csv_safe("") == ""


def test_sap_import_errors_xlsx_is_not_formula(client, admin_headers, db_session):
    # audyt SEC-008: kod/nazwa z wgranego pliku SAP szły do raportu błędów gołym sheet.append
    log = models.SapImport(kind="lfa1", filename="lfa1.xlsx", counts={},
                           errors=[{"row": 3, "sap_code": EVIL, "name": "=1+1",
                                    "reason": "brak nazwy"}])
    db_session.add(log)
    db_session.commit()
    report = client.get(f"/api/import/sap-imports/{log.id}/errors.xlsx", headers=admin_headers)
    assert report.status_code == 200
    row = _reload(report.content).active[2]
    assert (row[1].value, row[2].value) == (EVIL, "=1+1")
    assert row[1].data_type == "s" and row[2].data_type == "s"


def test_no_raw_worksheet_append_outside_exports():
    # strażnik SEC-008: wiersze arkuszy xlsx zapisujemy wyłącznie przez exports.append_row
    import pathlib
    import re
    app = pathlib.Path(__file__).resolve().parents[1] / "app"
    raw = re.compile(r"\b(ws|sheet|worksheet)\.append\(")
    offenders = [f"{path.relative_to(app)}:{no}"
                 for path in sorted(app.rglob("*.py")) if path.name != "exports.py"
                 for no, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1)
                 if raw.search(line)]
    assert offenders == []
