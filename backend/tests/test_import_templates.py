"""Wzory importów (2026-09-24): każdy wzór czyta się PRAWDZIWYM parserem swojego importu
i daje wszystkie kolumny ze wzoru — rozjazd nagłówków wzór↔parser = czerwony test."""
import io

import pytest
from openpyxl import load_workbook

from app.import_templates import TEMPLATES, build_template
from app.importers.excel import _map_headers
from app.importers.lfa1 import parse_lfa1
from app.importers.master_data import (
    _parse_cport_rows, _parse_dlt_stock_rows, _parse_marm_rows)
from app.importers.purchasing import EKKO_HEADERS, ETD_HEADERS, _parse_ekko_rows, _parse_etd_rows
from app.importers.queue import HEADER_MAP, _parse_rows
from app.invoices.master_import import load_xlsx_rows, parse_workbook_rows
from app.routers.imports import SAP_HEADERS
from app.tabular import normalize_header


def _headers(kind):
    return load_workbook(io.BytesIO(build_template(kind)))["Dane"][1]


def _all_mapped(kind, mapping):
    """Każdy nagłówek wzoru trafia w INNY klucz parsera (brak kolizji fragmentów)."""
    cells = [c.value for c in _headers(kind)]
    cols = _map_headers(cells, mapping)
    assert len(cols) == len(cells), (kind, sorted(cols))
    return cols


def test_every_template_downloads(client, admin_headers):
    for kind in TEMPLATES:
        r = client.get(f"/api/import/templates/{kind}", headers=admin_headers)
        assert r.status_code == 200, kind
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Dane", "Instrukcja"]   # dane PIERWSZE (parsery czytają arkusz 1)
    assert client.get("/api/import/templates/nie-ma", headers=admin_headers).status_code == 404


@pytest.mark.parametrize("kind,mapping", [("ekko", EKKO_HEADERS), ("etd", ETD_HEADERS),
                                          ("queue", HEADER_MAP), ("ref", SAP_HEADERS)])
def test_sap_style_headers_map_one_to_one(kind, mapping):
    _all_mapped(kind, mapping)


def test_parsers_read_example_rows():
    marm, _ = _parse_marm_rows(build_template("marm"))
    assert marm[0]["material_no"] == "REF12345" and marm[0]["unit"] == "KAR" and marm[0]["numerator"] == 12
    assert parse_lfa1(build_template("lfa1"))[0][0]["sap_code"] == "100200"
    assert _parse_dlt_stock_rows(build_template("dlt"))[0]["produkt"] == "REF12345"
    assert _parse_cport_rows(build_template("ports"), "wzor.xlsx")[0]["code"] == "CNNGB"
    ekko = _parse_ekko_rows(build_template("ekko"))[0]
    assert ekko["order_number"] == "4500617421" and ekko["delivery_date"] is not None   # data = komórka DATA
    assert _parse_etd_rows(build_template("etd"))[0]["etd"] is not None
    assert _parse_rows(build_template("queue"))[0]["container_no"] == "MSDU0806613"
    mat = parse_workbook_rows(load_xlsx_rows(build_template("materials")))[0]
    assert mat["ref_code"] == "REF12345" and mat["tariff_cn"] == "94017100"


@pytest.mark.parametrize("kind,required", [("po", {"container_no", "order_numbers"}),
                                           ("paz", {"produkt", "sztuk_na_palete"}),
                                           ("issues", {"data", "produkt", "ilosc"})])
def test_exact_header_imports(kind, required):
    assert required <= {normalize_header(c.value) for c in _headers(kind)}
