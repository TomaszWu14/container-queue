"""GET /api/invoice-batches/{id}/agency-mail.eml — szkic maila do agencji dla Outlooka."""
import email
from email.policy import default

from sqlalchemy import select

from app.models import AuditLog, Company, Container, CustomsAgency, InvoiceBatch
from tests.conftest import login
from tests.test_invoices_api import _company_id, _container, _import_master, _upload, fake_pdf  # noqa: F401


def _ready_batch(client, headers):
    """Kontener + paczka z zatwierdzoną fakturą i aktualnym Excelem."""
    _import_master(client, headers)
    cid = _container(client, headers)
    batch = _upload(client, headers, cid).json()
    job = batch["jobs"][0]
    items = client.get(f"/api/invoice-jobs/{job['id']}", headers=headers).json()["items"]
    body = {"items": [{"id": i["id"], "master_ref": "MSK2" if i["raw_ref"] == "MSK" else i["raw_ref"],
                       "qty": i["qty"], "amount": i["amount"], "weight_net": "", "weight_gross": "",
                       "skipped": i["raw_ref"] == "NOPE-1"} for i in items], "confirm": True, "conformity_reason": "test: brak danych SAP"}
    assert client.put(f"/api/invoice-jobs/{job['id']}/review", headers=headers, json=body).status_code == 200
    assert client.post(f"/api/invoice-batches/{batch['id']}/export", headers=headers).status_code == 200
    return cid, batch["id"]


def _assign_agency(db, cid, email_value="odprawy@agencja.pl; biuro@agencja.pl"):
    agency = CustomsAgency(name="Agencja Ćma", email=email_value)
    db.add(agency)
    db.flush()
    container = db.get(Container, cid)
    container.customs_agency_id = agency.id
    db.get(Company, container.company_id).avizo_cc = "zespol@firma.pl, ODPRAWY@agencja.pl"
    db.commit()
    return agency


def _url(batch_id):
    return f"/api/invoice-batches/{batch_id}/agency-mail.eml"


def test_eml_success(client, admin_headers, fake_pdf, db_session):
    cid, bid = _ready_batch(client, admin_headers)
    _assign_agency(db_session, cid)
    resp = client.get(_url(bid), headers=admin_headers)
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("message/rfc822")
    assert resp.headers["content-disposition"] == 'attachment; filename="faktury_MSDU0806613.eml"'
    msg = email.message_from_bytes(resp.content, policy=default)
    assert msg["X-Unsent"] == "1" and msg["From"] is None
    assert msg["To"] == "odprawy@agencja.pl, biuro@agencja.pl"
    assert msg["Cc"] == "zespol@firma.pl"   # adres z „Do” nie dubluje się w DW
    assert msg["Subject"].startswith("Faktury — MSDU0806613 — ")
    names = [a.get_filename() for a in msg.iter_attachments()]
    # Excel pozycji, kartoteka symboli (wzór agencji, 44 kolumny), PDF faktury
    # Excel, kartoteka, faktura i packing lista z zestawu (audyt #12 — komplet do odprawy)
    assert names[2:] == ["cipl_faktura_s1.pdf", "cipl_PL_s2.pdf"] and names[0].endswith(".xlsx")
    assert names[1] == "kartoteka_symboli_MSDU0806613.xlsx"
    import io
    from datetime import UTC, datetime

    import openpyxl
    # Excel paczki: czytelna nazwa z datą, MIME xlsx, treść = pozycje zatwierdzonej faktury
    excel = list(msg.iter_attachments())[0]
    assert names[0] == f"Faktury_MSDU0806613_{datetime.now(UTC).date().isoformat()}.xlsx"   # data UTC jak w bazie
    assert excel.get_content_type() == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    cells = {c for row in openpyxl.load_workbook(io.BytesIO(excel.get_content())).active.iter_rows(values_only=True)
             for c in row}
    assert "NL753-S-40" in cells and "NOPE-1" not in cells   # pominięta pozycja nie trafia do Excela
    assert list(msg.iter_attachments())[2].get_content_type() == "application/pdf"
    assert "zestawienie faktur w Excelu" in msg.get_body(("plain",)).get_content()
    part = list(msg.iter_attachments())[1]
    rows = list(openpyxl.load_workbook(io.BytesIO(part.get_content())).active.iter_rows(values_only=True))
    assert len(rows[0]) == 44 and rows[0][:3] == ("Symbol", "KodPCN", "KodTaric")
    by_symbol = {r[0]: r for r in rows[1:]}   # NOPE-1 pominięte, bez master — poza kartoteką
    assert sorted(by_symbol) == ["MSK2", "NL753-S-40"]
    cewnik = by_symbol["NL753-S-40"]
    assert cewnik[1:3] == ("9018", "00") and cewnik[11:13] == ("Cewnik", "Catheter")
    assert cewnik[14] == "SZT" and cewnik[23:27] == ("N", "N", "N", 106) and cewnik[27] == "admin"
    log = db_session.scalars(select(AuditLog).where(AuditLog.field == "invoice_agency_mail")).one()
    assert log.entity_id == cid and "szkic maila" in log.note


def test_eml_refusals(client, admin_headers, fake_pdf, db_session):
    cid, bid = _ready_batch(client, admin_headers)
    resp = client.get(_url(bid), headers=admin_headers)
    assert resp.status_code == 409 and "agencji" in resp.json()["detail"]
    agency = _assign_agency(db_session, cid, email_value="")
    resp = client.get(_url(bid), headers=admin_headers)
    assert resp.status_code == 409 and "e-mail" in resp.json()["detail"]
    agency.email = "odprawy@agencja.pl"
    batch = db_session.get(InvoiceBatch, bid)
    batch.ready = False   # Excel nieaktualny
    db_session.commit()
    resp = client.get(_url(bid), headers=admin_headers)
    assert resp.status_code == 409 and "Excel" in resp.json()["detail"]
    batch.ready = True
    db_session.commit()
    (fake_pdf / batch.jobs[0].stored_name).unlink()
    resp = client.get(_url(bid), headers=admin_headers)
    assert resp.status_code == 409 and "Brak pliku" in resp.json()["detail"]


def test_eml_unconfirmed_batch_blocked(client, admin_headers, fake_pdf, db_session):
    """Paczka bez zatwierdzonych faktur (Excel nigdy nie wygenerowany) — blokada z instrukcją."""
    _import_master(client, admin_headers)
    cid = _container(client, admin_headers)
    bid = _upload(client, admin_headers, cid).json()["id"]
    _assign_agency(db_session, cid)
    resp = client.get(_url(bid), headers=admin_headers)
    assert resp.status_code == 409 and "wygeneruj Excel" in resp.json()["detail"]
    assert db_session.scalars(select(AuditLog).where(AuditLog.field == "invoice_agency_mail")).first() is None


def test_eml_isolation_and_roles(client, admin_headers, fake_pdf, db_session):
    cid, bid = _ready_batch(client, admin_headers)
    _assign_agency(db_session, cid)
    other = _company_id(client, admin_headers, index=1)
    client.post("/api/users", headers=admin_headers, json={
        "login": "zak", "password": "haslo123", "role": "purchasing", "company_id": other})
    assert client.get(_url(bid), headers=login(client, "zak", "haslo123")).status_code == 404
    agency_id = db_session.get(Container, cid).customs_agency_id
    client.post("/api/users", headers=admin_headers, json={
        "login": "agent", "password": "haslo123", "role": "customs", "customs_agency_id": agency_id})
    assert client.get(_url(bid), headers=login(client, "agent", "haslo123")).status_code == 403


def test_cn_codes_taric_from_customs_code():
    from app.invoices.symbols import cn_codes
    assert cn_codes("4823 70 90", "") == ("48237090", "00")
    assert cn_codes("62101098", "6210109800") == ("62101098", "00")
    assert cn_codes("", "621010981234") == ("62101098", "12")


def test_symbols_follow_master_data_template(client, admin_headers, fake_pdf, db_session):
    """Kartoteka symboli w mailu = wzór z Master data → Wzory plików dla agencji: zmiana nazwy,
    kolejności, stałej i usunięcie kolumny w kafelku zmienia plik bez zmiany kodu."""
    import io

    import openpyxl
    cid, bid = _ready_batch(client, admin_headers)
    _assign_agency(db_session, cid)
    cols = client.get("/api/agency-templates", headers=admin_headers).json()["templates"][0]["columns"]
    cols = [c for c in cols if c["name"] != "KodUE1"]                     # agencja usunęła kolumnę
    cols[0], cols[1] = cols[1], cols[0]                                   # KodPCN przed Symbol
    next(c for c in cols if c["name"] == "Militarny").update(value="T")   # inna stała
    next(c for c in cols if c["name"] == "Uwagi").update(name="Notatka", source="const", value="z TIMPORYE")
    assert client.put("/api/agency-templates/winsad_symbole", headers=admin_headers,
                      json={"columns": cols}).status_code == 200

    msg = email.message_from_bytes(client.get(_url(bid), headers=admin_headers).content, policy=default)
    part = list(msg.iter_attachments())[1]
    rows = list(openpyxl.load_workbook(io.BytesIO(part.get_content())).active.iter_rows(values_only=True))
    head = rows[0]
    assert len(head) == 43 and head[:2] == ("KodPCN", "Symbol") and "KodUE1" not in head
    cewnik = next(dict(zip(head, r)) for r in rows[1:] if r[1] == "NL753-S-40")
    assert (cewnik["KodPCN"], cewnik["Militarny"], cewnik["Notatka"]) == ("9018", "T", "z TIMPORYE")
    assert cewnik["IDZestawu"] == 106 and cewnik["NazwaJednMiar"] == "SZT"


def test_preview_shows_exactly_what_eml_contains(client, admin_headers, fake_pdf, db_session):
    """Podgląd przed wysyłką (2026-10-06 „nie widzę, co wyszło”): odbiorcy, temat, treść,
    załączniki z rozmiarami = to samo co .eml; ostrzeżenia; kartoteka do pobrania; bez audytu."""
    import io

    import openpyxl
    cid, bid = _ready_batch(client, admin_headers)
    _assign_agency(db_session, cid)
    resp = client.get(f"/api/invoice-batches/{bid}/agency-mail/preview", headers=admin_headers)
    assert resp.status_code == 200, resp.text
    prev = resp.json()
    assert prev["to"] == ["odprawy@agencja.pl", "biuro@agencja.pl"] and prev["cc"] == ["zespol@firma.pl"]
    assert prev["agency"] == "Agencja Ćma" and "MSDU0806613" in prev["subject"]
    assert [a["kind"] for a in prev["attachments"]] == ["excel", "symbols", "invoice", "packing_list"]
    assert all(a["size"] > 0 for a in prev["attachments"]) and prev["attachments"][2]["job_id"]
    assert "Załączniki:" in prev["body"]
    # pozycja bez dopasowania do master (NOPE-1 pominięta → nie liczy się) — tu wszystko dopasowane
    assert not any("niezatwierdzone" in w for w in prev["warnings"])
    assert db_session.scalars(select(AuditLog).where(AuditLog.field == "invoice_agency_mail")).first() is None

    msg = email.message_from_bytes(client.get(_url(bid), headers=admin_headers).content, policy=default)
    assert [p.get_filename() for p in msg.iter_attachments()] == [a["name"] for a in prev["attachments"]]
    xlsx = client.get(f"/api/invoice-batches/{bid}/symbols.xlsx", headers=admin_headers)
    assert xlsx.status_code == 200 and "kartoteka_symboli_MSDU0806613" in xlsx.headers["content-disposition"]
    rows = list(openpyxl.load_workbook(io.BytesIO(xlsx.content)).active.iter_rows(values_only=True))
    assert sorted(r[0] for r in rows[1:]) == ["MSK2", "NL753-S-40"]


def test_preview_warns_about_required_template_column_empty(client, admin_headers, fake_pdf, db_session):
    cid, bid = _ready_batch(client, admin_headers)
    _assign_agency(db_session, cid)
    cols = client.get("/api/agency-templates", headers=admin_headers).json()["templates"][0]["columns"]
    next(c for c in cols if c["name"] == "DodOpisTowaru")["required"] = True   # kolumna zawsze pusta
    assert client.put("/api/agency-templates/winsad_symbole", headers=admin_headers,
                      json={"columns": cols}).status_code == 200
    prev = client.get(f"/api/invoice-batches/{bid}/agency-mail/preview", headers=admin_headers).json()
    assert any("„DodOpisTowaru” pusta w 2 wierszach" in w for w in prev["warnings"])


def test_mail_includes_bl_uploaded_as_container_document(client, admin_headers, fake_pdf, db_session):
    """B/L wgrany osobno (typ z kafelkiem BL) trafia do maila, gdy zestaw go nie ma."""
    import io

    from app.models import DocumentType
    from tests.conftest import pdf_bytes
    cid, bid = _ready_batch(client, admin_headers)
    _assign_agency(db_session, cid)
    bl = DocumentType(name="Konosament", is_active=True, tile_code="BL")
    db_session.add(bl)
    db_session.commit()
    up = client.post(f"/api/containers/{cid}/attachments", headers=admin_headers,
                     data={"document_type_id": str(bl.id)},
                     files={"file": ("BL STS1.pdf", io.BytesIO(pdf_bytes("bl")), "application/pdf")})
    assert up.status_code == 201, up.text
    prev = client.get(f"/api/invoice-batches/{bid}/agency-mail/preview", headers=admin_headers).json()
    bl_file = next(a for a in prev["attachments"] if a["kind"] == "bl")
    assert bl_file["name"] == "BL STS1.pdf" and bl_file["attachment_id"] == up.json()["id"]
    assert "konosament (B/L)" in prev["body"] and "packing listy" in prev["body"]



def test_partial_export_is_marked_in_subject_and_body(client, admin_headers, fake_pdf, db_session):
    """Audyt 2026-10-06 #16: Excel z części faktur nie wygląda w mailu jak komplet."""
    from app.models import InvoiceJob, InvoiceJobStatus
    cid, bid = _ready_batch(client, admin_headers)
    _assign_agency(db_session, cid)
    confirmed = db_session.query(InvoiceJob).filter_by(batch_id=bid, status=InvoiceJobStatus.confirmed).one()
    db_session.add(InvoiceJob(batch_id=bid, filename="druga.pdf", stored_name=confirmed.stored_name,
                              source_name=confirmed.source_name, status=InvoiceJobStatus.extracted))
    db_session.commit()
    prev = client.get(f"/api/invoice-batches/{bid}/agency-mail/preview", headers=admin_headers).json()
    assert prev["subject"].endswith("(częściowe 1/2)") and "Liczba faktur: 1 z 2" in prev["body"]


def test_mail_takes_newest_bl_loose_ci_and_marks_sent(client, admin_headers, fake_pdf, db_session):
    """§4 pkt 23, 24, 39: tylko najnowszy B/L z kafelka, CI wgrane jako plik też w mailu,
    typ MIME z rozszerzenia; szkic = WYSŁANE + ślad id plików w historii."""
    import io

    from app.models import DocumentStatus, DocumentType
    from tests.conftest import pdf_bytes
    cid, bid = _ready_batch(client, admin_headers)
    _assign_agency(db_session, cid)
    types = {code: DocumentType(name=code, is_active=True, tile_code=code) for code in ("BL", "CI")}
    db_session.add_all(types.values())
    db_session.commit()

    def up(code, name, tag, ctype="application/pdf"):
        r = client.post(f"/api/containers/{cid}/attachments", headers=admin_headers,
                        data={"document_type_id": str(types[code].id)},
                        files={"file": (name, io.BytesIO(pdf_bytes(tag)), ctype)})
        assert r.status_code == 201, r.text
        return r.json()["id"]

    up("BL", "BL stary.pdf", "bl1")
    new_bl = up("BL", "BL nowy.pdf", "bl2")
    ci = up("CI", "CI luzem.pdf", "ci", ctype="application/x-dziwny")
    prev = client.get(f"/api/invoice-batches/{bid}/agency-mail/preview", headers=admin_headers).json()
    assert [a["attachment_id"] for a in prev["attachments"] if a["kind"] == "bl"] == [new_bl]
    assert ci in [a["attachment_id"] for a in prev["attachments"] if a["kind"] == "invoice"]

    resp = client.get(_url(bid), headers=admin_headers)
    msg = email.message_from_bytes(resp.content, policy=default)
    loose = next(a for a in msg.iter_attachments() if a.get_filename() == "CI luzem.pdf")
    assert loose.get_content_type() == "application/pdf"
    db_session.expire_all()
    assert db_session.get(Container, cid).document_status == DocumentStatus.WYSLANE
    log = db_session.scalars(select(AuditLog).where(AuditLog.field == "invoice_agency_mail")).one()
    assert f"[plik {ci}]" in log.note and "[dok " in log.note
