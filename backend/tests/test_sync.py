import datetime

from app.database import SessionLocal
from app.models import Company, ContainerStatus, CustomsStatus
from app.routers.imports import build_container_fields


def _company(db):
    return db.scalar(__import__("sqlalchemy").select(Company).where(Company.code == "BOREALIS"))


def test_build_container_fields_maps_excel_row(client):
    db = SessionLocal()
    try:
        raw = {
            "supplier": "ACME", "vessel": "MAERSK", "eta": datetime.datetime(2026, 9, 1),
            "transport": "kolej", "container_no": "MSKU7026499", "order_numbers": "4500",
            "delivery_note": "", "purchase_note": "", "warehouse": "ACME",
            "incoming_delivery_no": "", "rf_number": "", "forwarder": "SPEDALFA",
            "document_flow": "", "customs_agency": "", "customs": "odprawiony",
            "notify_date": datetime.datetime(2026, 9, 2), "sent_required": "TAK",
            "sent_number": "", "sent_status": "",
        }
        fields = build_container_fields(db, _company(db), raw)
        assert fields["vessel"] == "MAERSK"
        assert fields["eta"] == datetime.date(2026, 9, 1)
        assert fields["customs_status"] == CustomsStatus.ODPRAWIONY
        assert fields["status"] == ContainerStatus.AWIZOWANY  # notify_date + odprawiony
        assert "company_id" not in fields and "transport_id" not in fields
    finally:
        db.rollback()
        db.close()


from sqlalchemy import select

from app.models import Container
from app.routers.imports import reconcile_queue


def _raw(no, eta_day, customs="", notify_day=None):
    return {
        "supplier": "ACME", "vessel": "MAERSK",
        "eta": datetime.datetime(2026, 9, eta_day),
        "transport": "kolej", "container_no": no, "order_numbers": "",
        "delivery_note": "", "purchase_note": "", "warehouse": "",
        "incoming_delivery_no": "", "rf_number": "", "forwarder": "",
        "document_flow": "", "customs_agency": "", "customs": customs,
        "notify_date": datetime.datetime(2026, 9, notify_day) if notify_day else None,
        "sent_required": "", "sent_number": "", "sent_status": "",
    }


import io

from openpyxl import Workbook

from app.models import AuditLog, Notification


def _xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["DOSTAWCA", "STATEK", "ETA", "KOLEJ/KOŁA", "NR KONTENERA",
               "STATUS ODPRAWY", "DATA ROZŁADUNKU"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_queue_sync_updates_notifies_and_audits(client, admin_headers):
    hdr = admin_headers
    base = "/api/import/queue-sync?company_code=BOREALIS&dry_run=false"
    no = "MSKU7026492"
    # 1. pierwszy snapshot — insert
    f1 = _xlsx([["ACME", "MAERSK", datetime.datetime(2026, 9, 1), "kolej", no, "", ""]])
    r1 = client.post(base, headers=hdr, files={"file": ("q.xlsx", f1)})
    assert r1.status_code == 200, r1.text
    assert r1.json()["changed"] == 1
    # 2. drugi snapshot — zmiana ETA → update + audyt + powiadomienie
    f2 = _xlsx([["ACME", "MAERSK", datetime.datetime(2026, 9, 9), "kolej", no, "", ""]])
    r2 = client.post(base, headers=hdr, files={"file": ("q.xlsx", f2)})
    assert r2.json()["changed"] == 1 and r2.json()["notified"] >= 1
    # 3. idempotencja — ten sam plik, 0 zmian
    r3 = client.post(base, headers=hdr, files={"file": ("q.xlsx", f2)})
    assert r3.json()["changed"] == 0
    # 4. dry_run (domyślnie true) liczy zmianę, ale nie zapisuje
    f4 = _xlsx([["ACME", "MAERSK", datetime.datetime(2026, 9, 20), "kolej", no, "", ""]])
    rd = client.post("/api/import/queue-sync?company_code=BOREALIS", headers=hdr,
                     files={"file": ("q.xlsx", f4)})
    assert rd.json()["dry_run"] is True and rd.json()["changed"] == 1
    # a realny stan nietknięty — kolejny dry-run nadal widzi tę samą zmianę
    rd2 = client.post("/api/import/queue-sync?company_code=BOREALIS", headers=hdr,
                      files={"file": ("q.xlsx", f4)})
    assert rd2.json()["changed"] == 1
    # audyt zawiera wpis o zmianie eta
    from app.database import SessionLocal
    db = SessionLocal()
    try:
        assert db.query(AuditLog).filter(AuditLog.field == "eta",
                                         AuditLog.note == "sync z Excela").count() >= 1
        assert db.query(Notification).filter(Notification.kind == "excel-sync").count() >= 1
    finally:
        db.close()


def _xlsx_sheets(sheets: dict):
    """Skoroszyt wielozakładkowy: {nazwa_zakładki: [wiersze]}. Pierwsza zakładka aktywna."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        ws.append(["DOSTAWCA", "STATEK", "ETA", "KOLEJ/KOŁA", "NR KONTENERA",
                   "STATUS ODPRAWY", "DATA ROZŁADUNKU"])
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_sync_sheet_param_wins_over_company_heuristic(client, admin_headers):
    """Jawny `sheet` (np. „2026") ma pierwszeństwo przed heurystyką po kodzie spółki.
    Bez niego _pick_sheet wybrałby zakładkę z kodem spółki w tytule."""
    hdr = admin_headers
    stary, nowy = "MSKU7027586", "MSKU7026595"
    content = _xlsx_sheets({
        "BOREALIS 2025": [["ACME", "MAERSK", datetime.datetime(2026, 9, 1), "kolej", stary, "", ""]],
        "2026": [["ACME", "MAERSK", datetime.datetime(2026, 9, 2), "kolej", nowy, "", ""]],
    })
    r = client.post("/api/import/queue-sync?company_code=BOREALIS&sheet=2026&dry_run=false",
                    headers=hdr, files={"file": ("q.xlsx", content)})
    assert r.status_code == 200, r.text
    assert r.json()["containers"] == 1

    db = SessionLocal()
    try:
        company = _company(db)
        got = db.scalars(select(Container.container_no).where(
            Container.company_id == company.id,
            Container.container_no.in_([stary, nowy]))).all()
        assert nowy in got, "nie wczytano zakładki wskazanej przez ?sheet="
        assert stary not in got, "wczytano zakładkę z kodem spółki mimo jawnego ?sheet="
    finally:
        db.close()


def test_sync_unknown_sheet_falls_back_to_company_heuristic(client, admin_headers):
    """Literówka/zmiana nazwy zakładki nie może wywalić syncu — fallback na dotychczasową
    heurystykę po kodzie spółki."""
    no = "MSKU4308321"
    content = _xlsx_sheets({
        "BOREALIS 2025": [["ACME", "MAERSK", datetime.datetime(2026, 9, 3), "kolej", no, "", ""]],
    })
    r = client.post("/api/import/queue-sync?company_code=BOREALIS&sheet=NIE-MA-TAKIEJ&dry_run=false",
                    headers=admin_headers, files={"file": ("q.xlsx", content)})
    assert r.status_code == 200, r.text
    assert r.json()["containers"] == 1


def test_queue_sync_requires_auth(client):
    """Bez sesji (rola editor) upload kolejki jest odrzucany — dawny token n8n zniknął."""
    r = client.post("/api/import/queue-sync?company_code=BOREALIS",
                    files={"file": ("q.xlsx", b"x")})
    assert r.status_code == 401


def test_sync_malformed_upload_is_not_500(client, admin_headers):
    r = client.post("/api/import/queue-sync?company_code=BOREALIS&dry_run=false",
                    headers=admin_headers, files={"file": ("q.xlsx", b"notxlsx")})
    assert r.status_code in (400, 422)


def test_sync_unknown_company_404(client, admin_headers):
    r = client.post("/api/import/queue-sync?company_code=NIEISTNIEJE&dry_run=false",
                    headers=admin_headers, files={"file": ("q.xlsx", _xlsx([]))})
    assert r.status_code == 404


def test_reconcile_status_only_change_is_silent(client):
    db = SessionLocal()
    try:
        company = _company(db)
        no = "MSKU7026492"
        r1 = reconcile_queue(db, company, [_raw(no, 1)], None)
        db.flush()
        assert r1[0]["changes"][0][0] == "__new__"
        cont = db.scalar(select(Container).where(Container.company_id == company.id,
                                                 Container.container_no == no))
        derived = cont.status
        # narzucamy status niezgodny z tym, co wyliczy arkusz — jedyna "zmiana" to status
        other = (ContainerStatus.ZAPOWIEDZIANY if derived != ContainerStatus.ZAPOWIEDZIANY
                 else ContainerStatus.W_TRANSPORCIE)
        cont.status = other
        db.flush()
        r2 = reconcile_queue(db, company, [_raw(no, 1)], None)
        db.flush()
        assert r2 == []  # brak wpisu do results → brak powiadomienia
        # three-way: status w apce != baseline, Excel == baseline -> "tylko apka" -> pending,
        # apka wygrywa lokalnie (status NIE jest nadpisywany przez Excel)
        assert cont.status == other
        after_audit = db.query(AuditLog).filter(AuditLog.entity_id == cont.id,
                                                 AuditLog.note == "sync z Excela",
                                                 AuditLog.field == "status").count()
        assert after_audit == 0
    finally:
        db.rollback()
        db.close()


import datetime as _dt

from app.models import ContainerStatus as _CS
from app.routers.imports import SYNCED_FIELDS, container_baseline, ser


def test_ser_normalizes_types():
    assert ser(None) is None
    assert ser(_CS.W_PORCIE) == _CS.W_PORCIE.value
    assert ser(_dt.date(2026, 9, 1)) == "2026-09-01"
    assert ser(5) == "5"
    assert ser("x") == "x"


def test_container_baseline_covers_synced_fields(client):
    db = SessionLocal()
    try:
        c = db.scalar(select(Container)) or None
        if c is None:
            import pytest; pytest.skip("brak kontenera w bootstrapie")
        b = container_baseline(c)
        assert set(b.keys()) == set(SYNCED_FIELDS)
    finally:
        db.close()


def test_reconcile_insert_update_idempotent(client):
    db = SessionLocal()
    try:
        company = _company(db)
        no = "MSKU7026492"
        # 1. insert
        r1 = reconcile_queue(db, company, [_raw(no, 1)], None)
        db.flush()
        assert r1[0]["changes"][0][0] == "__new__"
        cont = db.scalar(select(Container).where(Container.company_id == company.id,
                                                 Container.container_no == no))
        assert cont.eta == datetime.date(2026, 9, 1)
        # 2. update — zmiana ETA
        r2 = reconcile_queue(db, company, [_raw(no, 5)], None)
        db.flush()
        changed = {f for f, _, _ in r2[0]["changes"]}
        assert "eta" in changed
        assert cont.eta == datetime.date(2026, 9, 5)
        # 3. idempotencja — ten sam arkusz, 0 zmian
        r3 = reconcile_queue(db, company, [_raw(no, 5)], None)
        assert r3 == []
    finally:
        db.rollback()
        db.close()


import datetime as dt2

from sqlalchemy import select as _sel


def _rawB(no, vessel="MAERSK", eta_day=1):
    return {"supplier": "", "vessel": vessel, "eta": dt2.datetime(2026, 9, eta_day),
            "transport": "kolej", "container_no": no, "order_numbers": "", "delivery_note": "",
            "purchase_note": "", "warehouse": "", "incoming_delivery_no": "", "rf_number": "",
            "forwarder": "", "document_flow": "", "customs_agency": "", "customs": "",
            "notify_date": None, "sent_required": "", "sent_number": "", "sent_status": ""}


def test_bootstrap_baseline_none_excel_wins(client):
    db = SessionLocal()
    try:
        from app.models import Company
        company = db.scalar(_sel(Company).where(Company.code == "BOREALIS"))
        no = "MSKU7026492"
        reconcile_queue(db, company, [_rawB(no, "AAA")], None); db.flush()
        c = db.scalar(_sel(Container).where(Container.container_no == no,
                                            Container.company_id == company.id))
        assert c.sync_baseline is not None            # baseline zainicjowany
        assert c.sync_baseline["vessel"] == "AAA"     # B = E po bootstrapie
    finally:
        db.rollback(); db.close()


def test_app_change_goes_to_pending_not_excel(client):
    db = SessionLocal()
    try:
        from app.models import Company
        company = db.scalar(_sel(Company).where(Company.code == "BOREALIS"))
        no = "MSKU7026492"
        reconcile_queue(db, company, [_rawB(no, "AAA")], None); db.flush()
        c = db.scalar(_sel(Container).where(Container.container_no == no,
                                            Container.company_id == company.id))
        # zmiana po stronie apki (A != B), Excel bez zmian (E == B)
        c.vessel = "APP-EDIT"; db.flush()
        reconcile_queue(db, company, [_rawB(no, "AAA")], None); db.flush()
        # apka wygrywa: pole zostaje APP-EDIT, baseline NIE podniesiony (czeka na /applied)
        assert c.vessel == "APP-EDIT"
        assert c.sync_baseline["vessel"] == "AAA"     # B nietknięty -> trafi do pending
    finally:
        db.rollback(); db.close()


def test_excel_change_adopted_and_baseline_advances(client):
    db = SessionLocal()
    try:
        from app.models import Company
        company = db.scalar(_sel(Company).where(Company.code == "BOREALIS"))
        no = "MSKU7026492"
        reconcile_queue(db, company, [_rawB(no, "AAA")], None); db.flush()
        c = db.scalar(_sel(Container).where(Container.container_no == no,
                                            Container.company_id == company.id))
        reconcile_queue(db, company, [_rawB(no, "BBB")], None); db.flush()  # Excel zmienił
        assert c.vessel == "BBB"                      # A <- E
        assert c.sync_baseline["vessel"] == "BBB"     # B <- E
    finally:
        db.rollback(); db.close()


def test_conflict_app_newer_wins_writable_field(client):
    """Konflikt na polu writable (vessel): apka zmieniła PO ostatnim sync file_mtime -> apka wygrywa."""
    from app.audit import record
    from app.models import Company
    db = SessionLocal()
    try:
        company = db.scalar(_sel(Company).where(Company.code == "BOREALIS"))
        no = "MSKU7026492"
        reconcile_queue(db, company, [_rawB(no, "AAA")], None); db.flush()
        c = db.scalar(_sel(Container).where(Container.container_no == no,
                                            Container.company_id == company.id))
        old_before = c.updated_at
        c.vessel = "APP-EDIT"; db.flush()
        record(db, entity_type="containers", entity_id=c.id, field="vessel",
               old_value="AAA", new_value="APP-EDIT", user=None, note="test")
        db.flush()
        older_mtime = old_before - dt2.timedelta(hours=1)
        res = reconcile_queue(db, company, [_rawB(no, "BBB")], None, file_mtime=older_mtime)
        db.flush()
        assert c.vessel == "APP-EDIT"                 # apka wygrywa
        assert res == []
        assert c.sync_baseline["vessel"] == "AAA"     # B nietknięty -> pending
    finally:
        db.rollback(); db.close()


def test_conflict_excel_newer_wins_writable_field(client):
    """Konflikt na polu writable (vessel): plik nowszy niż zmiana apki -> Excel wygrywa."""
    from app.models import Company
    db = SessionLocal()
    try:
        company = db.scalar(_sel(Company).where(Company.code == "BOREALIS"))
        no = "MSKU7026492"
        reconcile_queue(db, company, [_rawB(no, "AAA")], None); db.flush()
        c = db.scalar(_sel(Container).where(Container.container_no == no,
                                            Container.company_id == company.id))
        c.vessel = "APP-EDIT"; db.flush()
        newer_mtime = dt2.datetime.utcnow() + dt2.timedelta(hours=1)
        res = reconcile_queue(db, company, [_rawB(no, "BBB")], None, file_mtime=newer_mtime)
        db.flush()
        assert c.vessel == "BBB"                      # Excel wygrywa
        assert c.sync_baseline["vessel"] == "BBB"
        changed = {f for f, _, _ in res[0]["changes"]}
        assert "vessel" in changed
    finally:
        db.rollback(); db.close()


def test_nonwritable_status_app_local_accepted(client):
    """Pole niezapisywalne (status): zmiana app-local (A!=B, E==B) jest akceptowana bez
    powiadomienia; B zostaje = E, żeby kolejny sync z tym samym Excelem jej nie cofnął."""
    db = SessionLocal()
    try:
        from app.models import Company
        company = db.scalar(_sel(Company).where(Company.code == "BOREALIS"))
        no = "MSKU7026492"
        r1 = reconcile_queue(db, company, [_raw(no, 1)], None)
        db.flush()
        assert r1[0]["changes"][0][0] == "__new__"
        c = db.scalar(select(Container).where(Container.company_id == company.id,
                                              Container.container_no == no))
        other = (ContainerStatus.ZAPOWIEDZIANY if c.status != ContainerStatus.ZAPOWIEDZIANY
                 else ContainerStatus.W_TRANSPORCIE)
        c.status = other
        db.flush()
        res = reconcile_queue(db, company, [_raw(no, 1)], None)
        db.flush()
        assert res == []                               # brak powiadomienia
        assert c.status == other                        # app-local zaakceptowany
        assert c.sync_baseline["status"] != other.value  # B = E (nie B <- A)
        assert reconcile_queue(db, company, [_raw(no, 1)], None) == []
        assert c.status == other                        # drugi sync też nie cofa
    finally:
        db.rollback(); db.close()


def test_writable_fields_exclude_status_and_notes(client):
    """WRITABLE_FIELDS = pola, w których edycja w apce wygrywa przy konflikcie (LWW).
    status/notes/customs_status muszą być Excel->apka only (Excel autorytatywny)."""
    from app.routers.imports import WRITABLE_FIELDS
    assert "status" not in WRITABLE_FIELDS and "notes" not in WRITABLE_FIELDS
    # customs_status jest lossy w Excelu (7 stanów apki -> 2) — nieautorytatywny w apce
    assert "customs_status" not in WRITABLE_FIELDS


def test_queue_sync_eta_stable_with_typed_date_cell(client, admin_headers):
    """Round-trip stabilny: ta sama eta zapisana raz jako datetime, raz jako date,
    nie generuje różnicy przy ponownym uploadzie."""
    hdr = admin_headers
    base = "/api/import/queue-sync?company_code=BOREALIS&dry_run=false"
    no = "MSKU7026532"
    f1 = _xlsx([["ACME", "MAERSK", datetime.datetime(2026, 9, 1), "kolej", no, "", ""]])
    r1 = client.post(base, headers=hdr, files={"file": ("q.xlsx", f1)})
    assert r1.status_code == 200 and r1.json()["changed"] == 1

    # ta sama eta, ale jako realna datetime.date w komórce
    f2 = _xlsx([["ACME", "MAERSK", datetime.date(2026, 9, 1), "kolej", no, "", ""]])
    r2 = client.post(base, headers=hdr, files={"file": ("q.xlsx", f2)})
    assert r2.status_code == 200, r2.text
    assert r2.json()["changed"] == 0   # brak różnicy -> round-trip stabilny
