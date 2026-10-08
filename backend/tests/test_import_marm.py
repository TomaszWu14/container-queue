import io

from openpyxl import Workbook
from sqlalchemy import select

from app.models import AuditLog, MaterialUnit
from app.routers.imports import _parse_marm_rows
from app.serializers import marm_volume_m3

# nagłówek jak w polskim eksporcie MARM (jednostka objętości zawiera fragment „OBJĘTOŚĆ” —
# mapowanie musi trafić w obie kolumny, nie w pierwszą pasującą)
HEADER = ['Materiał', 'Jedn. miary', 'Licznik', 'Mianownik',
          'Objętość', 'Jednostka objętości', 'Waga brutto', 'Jedn. wagi']


def _xlsx(rows, header=HEADER):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parse_marm_maps_columns_and_filters_placeholders():
    rows, placeholders = _parse_marm_rows(_xlsx([
        ['100200', 'SZT', 1, 1, 0, '0', 0.5, 'KG'],       # objętość 0 -> None, wiersz zostaje
        ['100200', 'KAR', 10, 1, 120, 'CDM', 5, 'KG'],
        ['100200', '0', 0, 0, 0, '0', 0, '0'],            # placeholder — odfiltrowany
    ]))
    assert placeholders == 1
    assert [r["unit"] for r in rows] == ["SZT", "KAR"]
    kar = rows[1]
    assert kar["material_no"] == "100200"
    assert kar["numerator"] == 10 and kar["denominator"] == 1
    assert kar["volume"] == 120 and kar["volume_unit"] == "CDM"
    assert kar["gross_weight"] == 5 and kar["weight_unit"] == "KG"
    assert rows[0]["volume"] is None                       # „0” = brak danych


def test_parse_marm_sap_headers():
    rows, _ = _parse_marm_rows(_xlsx(
        [['100300', 'KAR', 24, 1, 0.06, 'M3', None, '']],
        header=['MATNR', 'MEINH', 'UMREZ', 'UMREN', 'VOLUM', 'VOLEH', 'BRGEW', 'GEWEI']))
    assert rows == [{"material_no": "100300", "unit": "KAR", "numerator": 24,
                     "denominator": 1, "volume": 0.06, "volume_unit": "M3",
                     "gross_weight": None, "weight_unit": "",
                     "length": None, "width": None, "height": None,
                     "dimension_unit": "", "ean": ""}]


# --- wymiary jednostki (LAENG/BREIT/HOEH + MEABM) ---

def test_parse_marm_dimensions_sap_and_polish_headers():
    # SAP-owe nagłówki
    rows, _ = _parse_marm_rows(_xlsx(
        [['100300', 'KAR', 600, 400, 300, 'MM']],
        header=['MATNR', 'MEINH', 'LAENG', 'BREIT', 'HOEH', 'MEABM']))
    assert rows[0]["length"] == 600 and rows[0]["width"] == 400
    assert rows[0]["height"] == 300 and rows[0]["dimension_unit"] == "MM"
    # polskie nagłówki — „Jedn. wymiaru” nie może trafić w kolumnę jednostki miary
    rows, _ = _parse_marm_rows(_xlsx(
        [['100301', 'KAR', 60, 40, 30, 'CM', 0]],
        header=['Materiał', 'Jedn. miary', 'Długość', 'Szerokość',
                'Wysokość', 'Jedn. wymiaru', 'Objętość']))
    assert rows[0]["unit"] == "KAR" and rows[0]["dimension_unit"] == "CM"
    assert rows[0]["length"] == 60
    # wymiary „0” = placeholder SAP -> None
    rows, _ = _parse_marm_rows(_xlsx(
        [['100302', 'SZT', 0, 0, 0, '0']],
        header=['MATNR', 'MEINH', 'LAENG', 'BREIT', 'HOEH', 'MEABM']))
    assert rows[0]["length"] is None and rows[0]["height"] is None


def test_dims_cm_normalization():
    from app.models import MaterialUnit

    def unit_row(**kw):
        return MaterialUnit(material_no="1", unit="KAR", **kw)

    assert unit_row(length=600, width=400, height=300, dimension_unit="MM") \
        .dims_cm() == (60.0, 40.0, 30.0)
    assert unit_row(length=60, width=40, height=30, dimension_unit="CM") \
        .dims_cm() == (60.0, 40.0, 30.0)
    assert unit_row(length=0.6, width=0.4, height=0.3, dimension_unit="M") \
        .dims_cm() == (60.0, 40.0, 30.0)
    assert unit_row(length=0.6, width=0.4, height=0.3, dimension_unit="MTR") \
        .dims_cm() == (60.0, 40.0, 30.0)
    # brak kompletu wymiarów albo nieznana jednostka -> None
    assert unit_row(length=60, width=40, height=None, dimension_unit="CM").dims_cm() is None
    assert unit_row(length=60, width=40, height=30, dimension_unit="XYZ").dims_cm() is None


def test_import_marm_stores_dimensions(client, admin_headers, db_session):
    content = _xlsx(
        [['100200', 'KAR', 600, 400, 300, 'MM']],
        header=['MATNR', 'MEINH', 'LAENG', 'BREIT', 'HOEH', 'MEABM'])
    done = _upload(client, admin_headers, content, False)
    assert done.status_code == 200, done.text
    stored, = db_session.scalars(select(MaterialUnit)).all()
    assert float(stored.length) == 600 and stored.dimension_unit == "MM"
    assert stored.dims_cm() == (60.0, 40.0, 30.0)


def _upload(client, headers, content, dry_run):
    return client.post(
        f"/api/import/material-units?dry_run={str(dry_run).lower()}",
        headers=headers, files={"file": ("MARM.xlsx", content, "application/vnd.ms-excel")})


def test_import_marm_dry_run_upsert_and_audit(client, admin_headers, db_session):
    content = _xlsx([
        ['100200', 'SZT', 1, 1, 0, '0', None, ''],
        ['100200', 'KAR', 10, 1, 120, 'CDM', 5, 'KG'],
        ['100200', 'KAR', 10, 1, 120, 'CDM', 5, 'KG'],    # duplikat w pliku
        ['100200', '0', 0, 0, 0, '0', None, ''],          # placeholder
    ])
    preview = _upload(client, admin_headers, content, True)
    assert preview.status_code == 200, preview.text
    assert preview.json()["counts"] == {"total": 4, "new": 2, "updated": 0,
                                        "placeholders": 1, "duplicate": 1,
                                        "invalid": 0, "disappeared": 0,
                                        # materiały do dopasowania faktur (marm_sync)
                                        "materials_new": 1, "conversions": 1}
    assert db_session.scalars(select(MaterialUnit)).all() == []   # dry-run nic nie zapisuje

    done = _upload(client, admin_headers, content, False)
    assert done.status_code == 200, done.text
    stored = db_session.scalars(select(MaterialUnit).order_by(MaterialUnit.unit)).all()
    assert [(m.unit, m.numerator) for m in stored] == [("KAR", 10), ("SZT", 1)]

    # ponowny import aktualizuje, nie duplikuje
    again = _upload(client, admin_headers, content, False)
    assert again.json()["counts"]["new"] == 0 and again.json()["counts"]["updated"] == 2
    assert len(db_session.scalars(select(MaterialUnit)).all()) == 2

    # wpis audytu jak przy innych importach
    logs = db_session.scalars(select(AuditLog).where(
        AuditLog.entity_type == "material_units")).all()
    assert logs and "MARM" in (logs[0].note or "")


def test_import_marm_needs_login(client):
    response = _upload(client, {}, _xlsx([['1', 'SZT', 1, 1, 1, 'M3', None, '']]), True)
    assert response.status_code == 401


def test_list_material_units_search(client, admin_headers):
    _upload(client, admin_headers, _xlsx([
        ['100200', 'KAR', 10, 1, 120, 'CDM', None, ''],
        ['555555', 'SZT', 1, 1, 0.001, 'M3', None, ''],
    ]), False)
    response = client.get("/api/material-units?q=1002", headers=admin_headers)
    assert response.status_code == 200
    body = response.json()
    assert [r["material_no"] for r in body] == ["100200"]
    assert body[0]["volume"] == 120 and body[0]["volume_unit"] == "CDM"


def test_list_material_units_search_ignores_case(client, admin_headers):
    """Na Postgresie LIKE rozróżnia wielkość liter (SQLite nie) — wyszukiwarka musi
    porównywać lower() z obu stron, żeby „abc” znalazło „ABC100” także na produkcji."""
    from sqlalchemy import event

    from app.database import engine
    _upload(client, admin_headers, _xlsx([['ABC100', 'KAR', 10, 1, 120, 'CDM', None, '']]), False)
    sql: list[str] = []

    def grab(_conn, _cur, statement, *_a):
        sql.append(statement)
    event.listen(engine, "before_cursor_execute", grab)
    try:
        response = client.get("/api/material-units?q=abc", headers=admin_headers)
    finally:
        event.remove(engine, "before_cursor_execute", grab)
    assert [r["material_no"] for r in response.json()] == ["ABC100"]
    query = next(s for s in sql if "material_units" in s and "LIKE" in s)
    assert "lower(" in query.lower(), query


# --- wyliczanie objętości z MARM ---

class _Row:
    def __init__(self, unit, numerator=1, denominator=1, volume=None, volume_unit=""):
        self.unit, self.numerator, self.denominator = unit, numerator, denominator
        self.volume, self.volume_unit = volume, volume_unit


def test_volume_direct_unit_m3():
    rows = [_Row("KAR", 24, 1, volume=0.06, volume_unit="M3")]
    assert marm_volume_m3(10, "KAR", rows) == 0.6


def test_volume_cdm_converted_to_m3():
    rows = [_Row("KAR", 10, 1, volume=120, volume_unit="CDM")]
    assert marm_volume_m3(2, "KAR", rows) == 0.24


def test_volume_base_unit_via_conversion():
    # pozycja w SZT (baza), objętość tylko na KAR: 48 szt = 2 kartony po 60 CDM
    rows = [_Row("SZT"), _Row("KAR", 24, 1, volume=60, volume_unit="CDM")]
    assert round(marm_volume_m3(48, "SZT", rows), 6) == 0.12


def test_volume_unknown_returns_none():
    assert marm_volume_m3(5, "SZT", [_Row("SZT")]) is None            # brak objętości
    assert marm_volume_m3(5, "SZT", [_Row("SZT", volume=1, volume_unit="XYZ")]) is None


def test_container_items_expose_computed_volume(client, admin_headers, db_session):
    _upload(client, admin_headers, _xlsx([
        ['100200', 'KAR', 1, 1, 100, 'CDM', None, ''],
    ]), False)
    companies = client.get("/api/companies", headers=admin_headers).json()
    acme_id = next(c["id"] for c in companies if c["code"] == "ACME")
    created = client.post("/api/containers", headers=admin_headers, json={
        "container_no": "MSKU5696108", "company_id": acme_id})
    assert created.status_code in (200, 201), created.text
    container_id = created.json()["id"]
    from app.models import Container
    db_session.query(Container).filter(Container.id == container_id) \
        .update({"order_numbers": "4500011111"})
    db_session.commit()

    items_wb = Workbook()
    sheet = items_wb.active
    sheet.append(['Dok.zaopatrz.', 'Pozycja', 'Krótki tekst', 'Materiał', 'Ilość zam.',
                  'Jedn. miary', 'Waga netto', 'Waga brutto', 'Objętość',
                  'Jednostka obj.', 'Planowana data wysyłki'])
    sheet.append(['4500011111', '10', 'Rękawice', '100200', 3, 'KAR',
                  1, 2, 0.5, 'M3', None])
    buffer = io.BytesIO()
    items_wb.save(buffer)
    imported = client.post(
        "/api/import/order-items?company_code=ACME&dry_run=false",
        headers=admin_headers,
        files={"file": ("REF.xlsx", buffer.getvalue(), "application/vnd.ms-excel")})
    assert imported.status_code == 200, imported.text

    items = client.get(f"/api/containers/{container_id}/items", headers=admin_headers)
    assert items.status_code == 200, items.text
    item, = items.json()
    assert round(item["computed_volume_m3"], 6) == 0.3    # 3 KAR × 100 CDM
    assert item["volume"] == "0.5"                        # pole z pliku bez zmian (fallback UI)
