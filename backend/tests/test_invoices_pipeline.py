"""Logika modułu Faktury → Excel bez HTTP: liczby, cięcie, kolumny, wagi, master data."""
from decimal import Decimal

import pytest

from app.invoices import extractor, master_import, packing_list, splitter, uom
from app.invoices.numbers import normalize_number, number_text
from app.models import InvoiceDocKind

# realne nagłówki z próbki: podwójna spacja, literówki PROFOMA/PORFORMA, PL zawiera "invoice"
CI = "ACME CO COMMERCIAL  INVOICE Shipper No. & date of Invoice A1 2026"
PL = "ACME CO PACKING LIST Shipper No. & date of invoice A1 2026 CTNS"
PROF = "ACME CO PROFOMA  INVOICE No. & date of Invoice A1-S 2026"
BL = "DEMO MEDICAL BILL OF LADING DMBL4815162 SHANGHAI GDANSK POLAND"
CONT = "RNBL10001 Disposable gloves 987654321N 4,400 19.000 83,600.00 page 2"


@pytest.mark.parametrize("raw,expected", [
    ("13,500.00", Decimal("13500.00")), ("13.500,00", Decimal("13500.00")),
    ("13 500,00", Decimal("13500.00")), ("1.234.567,89", Decimal("1234567.89")),
    # kropka + 3 cyfry = DZIESIĘTNA (waga 12.500 kg); przecinek + 3 cyfry = angielskie tysiące
    ("2,500", Decimal("2500")), ("0,500", Decimal("0.500")), ("1.500", Decimal("1.500")),
    ("12.500", Decimal("12.500")), ("1,000.00", Decimal("1000.00")), ("1.000,00", Decimal("1000.00")),
    ("$ 250.5", Decimal("250.5")), ("1E3", Decimal("1E3")), ("1,2-3,4", None),
    # kody walut i jednostki doklejone w komórce
    ("250.00 USD", Decimal("250.00")), ("1,000 PCS", Decimal("1000")), ("12,5 kg", Decimal("12.5")),
    ("n/a", None), ("", None), (None, None),
    # prefiks kraju przy $ (faktury FICTIVA: „US$7,416.90” zostawało tekstem, suma = 0)
    ("US$7,416.90", Decimal("7416.90")), ("HK$ 12.00", Decimal("12.00")),
])
def test_normalize_number(raw, expected):
    assert normalize_number(raw) == expected


@pytest.mark.parametrize("value,expected", [
    # kody ze spacją (master: AT-SGS-XL_1) wypadały z pozycji po cichu
    ("AT-SGS-XL 1", True), ("AT-SD-S 1 BLUE-CN", True), ("AT-S-UNI7-C-CN", True),
    ("Disposable Gowns", False), ("BANK OF CHINA", False), ("TOTAL USD", False),
])
def test_is_ref_value_with_space(value, expected):
    from app.invoices.extractor import is_ref_value
    assert is_ref_value(value) is expected


def test_number_text_keeps_value_not_format():
    assert number_text("13.500,00") == "13500"
    assert number_text("13,500.00") == "13500"
    # kolumna wagi/ceny: „12,500” to 12.5, nie 12 500
    assert normalize_number("12,500", prefer_decimal=True) == Decimal("12.500")
    assert normalize_number("12,500") == Decimal("12500")
    assert number_text("12,5") == "12.5"
    assert number_text("abc") == ""


def test_classify_variants():
    assert splitter._marker_kind(CI) == InvoiceDocKind.invoice
    assert splitter._marker_kind(PL) == InvoiceDocKind.packing_list
    assert splitter._marker_kind(PROF) == InvoiceDocKind.proforma
    assert splitter._marker_kind(BL) == InvoiceDocKind.other
    assert splitter._marker_kind("DRAFT") is None and splitter._is_blank("DRAFT")   # skan/pusta
    assert splitter._marker_kind(CONT) is None and not splitter._is_blank(CONT)     # kontynuacja
    assert splitter.classify_first_page("DRAFT") == InvoiceDocKind.other


def test_split_groups_and_cuts(monkeypatch, tmp_path):
    written = []
    monkeypatch.setattr(splitter, "page_texts", lambda p: [CI, CONT, PL, PROF])
    monkeypatch.setattr(splitter, "write_range", lambda p, a, b, out: written.append((a, b, out)))
    parts = splitter.split_pdf(str(tmp_path / "zestaw.pdf"))
    assert [(p["kind"], p["page_from"], p["page_to"]) for p in parts] == [
        (InvoiceDocKind.invoice, 1, 2), (InvoiceDocKind.packing_list, 3, 3),
        (InvoiceDocKind.proforma, 4, 4)]
    assert len(written) == 3
    assert parts[0]["out_path"].endswith("zestaw_doc1_invoice.pdf")


def test_single_doc_returns_original_without_copy(monkeypatch):
    monkeypatch.setattr(splitter, "page_texts", lambda p: [PROF])
    monkeypatch.setattr(splitter, "write_range",
                        lambda *a: pytest.fail("nie powinno ciąć jednodokumentowego PDF"))
    assert splitter.split_pdf("x/proforma.pdf") == [
        {"kind": InvoiceDocKind.proforma, "page_from": 1, "page_to": 1, "out_path": "x/proforma.pdf",
         "text": PROF}]


def test_split_real_pdf_roundtrip(tmp_path):
    """pypdf faktycznie zapisuje zakres stron (bez monkeypatcha)."""
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(width=200, height=200)
    src = tmp_path / "trzy.pdf"
    with open(src, "wb") as handle:
        writer.write(handle)
    out = tmp_path / "dwie.pdf"
    splitter.write_range(str(src), 2, 3, str(out))
    assert len(PdfReader(str(out)).pages) == 2
    assert splitter.page_texts(str(src)) == ["", "", ""]


def test_identify_columns_multilingual():
    ci = extractor.identify_columns(["No.", "Item No.", "Description of goods", "Q'ty\n(pcs)",
                                     "Unit price", "Amount (USD)", "N.W. (kg)", "G.W. (kg)", "CTNS"])
    assert ci["ref"] == 1 and ci["desc"] == 2 and ci["qty"] == 3
    assert ci["price"] == 4 and ci["net"] == 5
    assert ci["weight_net"] == 6 and ci["weight_gross"] == 7 and ci["cartons"] == 8


def test_identify_columns_ref_desc_combined_and_sno():
    ci = extractor.identify_columns(["REF / Description", "Qty", "Net value"])
    assert ci["_ref_desc_combined"] == 0
    ci = extractor.identify_columns(["S.No", "Product", "Qty(pcs)", "Amount"])
    assert ci["ref"] == 0 and ci.get("_sno_format")


def test_column_map_text_and_apply():
    cmap = extractor.parse_column_map("ref=Kod towaru; qty = Sztuk ;net=Wartość; bogus=x\nunit=jm")
    assert cmap == {"ref": "Kod towaru", "qty": "Sztuk", "net": "Wartość", "unit": "jm"}
    ci = extractor.apply_column_map(["Lp", "Kod towaru", "Nazwa", "Sztuk", "jm", "Wartość"], cmap)
    assert ci == {"ref": 1, "qty": 3, "unit": 4, "net": 5}
    assert extractor.apply_column_map(["a", "b"], cmap) is None


def _pages():
    table = [
        ["No", "Item No.", "Description", "Qty", "Unit", "Unit price", "Amount"],
        ["1", "NL753-S-40", "Catheter", "1,000", "PCS", "0.25", "250.00"],
        ["2", "MSK-2", "Mask", "20", "CTN", "10", "200.00"],
        ["", "", "", "", "", "", ""],
        ["", "TOTAL", "", "1,020", "", "", "450.00"],
    ]
    text = ("COMMERCIAL INVOICE\nInvoice No.: FV/2026/9\nCONTAINER No.:MSNU8152864\n"
            "Terms of Delivery:\nFOB QINGDAO")
    return [(text, [table])]


def test_parse_pages_items_totals_meta():
    doc = extractor.parse_pages(_pages())
    assert [i["ref"] for i in doc.items] == ["NL753-S-40", "MSK-2"]
    assert doc.items[0]["qty"] == "1,000" and doc.items[0]["net"] == "250" and doc.items[0]["unit"] == "PCS"
    assert doc.total_net == "450.00" and doc.total_qty == "1,020"
    assert doc.invoice_number == "FV/2026/9"
    meta = extractor.header_meta(doc.raw_text)
    assert meta == {"container_no": "MSNU8152864", "delivery_terms": "FOB QINGDAO"}


def test_header_meta_rejects_non_incoterm_and_missing():
    assert extractor.header_meta("Terms of Delivery:\n\nInvoice No: INV12345")["delivery_terms"] == ""
    assert extractor.header_meta("brak danych") == {"container_no": "", "delivery_terms": ""}


def test_invoice_number_variants():
    assert extractor.extract_invoice_number("PROFORMA INVOICE NO. PI-2026-001 dated") == "PI-2026-001"
    assert extractor.extract_invoice_number("No. & date of invoice HS260917 2026-09-17") == "HS260917"
    assert extractor.extract_invoice_number("Invoice No.: DATE") == ""


def test_weight_map_sums_per_ref_and_skips_empty():
    items = [{"ref": "NL753-S-40", "weight_net": "12.5", "weight_gross": "13.2", "cartons": "3"},
             {"ref": "nl753 s 40", "weight_net": "1", "weight_gross": "", "cartons": "1"},
             {"ref": "EMPTY", "weight_net": "", "weight_gross": "", "cartons": ""}]
    wm = packing_list.weight_map_from_items(items)
    assert wm == {"NL753S40": {"weight_net": "13.5", "weight_gross": "13.2", "cartons": "4"}}


def test_uom_canonical_and_factor():
    assert uom.canonical_unit("Kartony") == "CTN" and uom.canonical_unit("szt.") == "PCS"
    assert uom.canonical_unit("pieces") == "PCS" and uom.canonical_unit("") == ""
    rules = [{"ref_norm": "X", "unit_from": "CTN", "unit_to": "PCS", "factor": 24.0},
             {"ref_norm": "*", "unit_from": "PAL", "unit_to": "PCS", "factor": 1000.0}]
    assert uom.get_factor("CTN", "PCS", "X", rules) == 24.0
    assert uom.get_factor("PCS", "CTN", "X", rules) == pytest.approx(1 / 24)
    assert uom.get_factor("PAL", "szt", "Y", rules) == 1000.0
    assert uom.get_factor("CTN", "PCS", "Y", rules) is None
    assert uom.get_factor("PCS", "PCS", "Y", rules) == 1.0


def test_master_two_row_header_with_converters():
    rows = [
        ["", "", "", "", "", "SZT", "SZT", "OP", "OP", "KAR", "KAR", "PAZ", "PPA"],
        ["ref_code", "opis_pl", "opis_en", "Podstawowa jednostka miary", "kod producenta",
         "sztuka_ean", "Ilość podstawowej jednostki miary",
         "op_ean", "Ilość podstawowej jednostki miary",
         "karton_ean", "Ilość podstawowej jednostki miary",
         "Ilość podstawowej jednostki miary PAZ", "Ilość podstawowej jednostki miary PPA"],
        ["R-1", "Opis", "Desc", "SZT", "PROD-9",
         "5907996800810", "1", "5907996800018", "12", "5907996800025", "240", "4800", "2400"],
        ["", "pusty wiersz bez REF"],
    ]
    recs = master_import.parse_workbook_rows(rows)
    assert len(recs) == 1
    rec = recs[0]
    assert rec["ref_code"] == "R-1" and rec["ref_norm"] == "R1" and rec["base_uom"] == "SZT"
    assert rec["producer_code"] == "PROD-9" and rec["sent"] is None
    assert rec["levels"]["op"]["qty_base"] == "12" and rec["levels"]["paz"]["qty_base"] == "4800"
    rules = master_import._conversion_rules(rec)
    # SZT→SZT (poziom bazowy) pominięty; reszta na kanoniczne jednostki
    assert ("PCS", "PCS", 1.0) not in rules
    assert ("OP", "PCS", 12.0) in rules and ("CTN", "PCS", 240.0) in rules
    assert ("PAZ", "PCS", 4800.0) in rules


def test_master_single_row_header_pl_names_and_flags():
    rows = [["REF", "Opis PL", "Kod CN", "SENT", "Stawka VAT", "EAN", "JM"],
            ["NL753-S-40", "Cewnik", "9018 39 00", "TAK", "23%", "5907996800810", "szt"],
            ["Z9", "Bez sentu", "", "", "8,0", "12345", "PCS"]]
    a, b = master_import.parse_workbook_rows(rows)
    assert a["name_pl"] == "Cewnik" and a["tariff_cn"] == "9018 39 00" and a["sent"] is True
    assert a["vat_rate"] == "23" and a["ean"] == "5907996800810"
    assert b["sent"] is None and b["vat_rate"] == "8" and b["ean"] == ""   # zły EAN → puste


# --- poprawki po dogłębnym code review ------------------------------------------------

def test_header_found_below_address_blocks_in_single_table_page():
    """Skan/PDF bez linii tabeli = jedna tabela na stronę; nagłówek pod 15 liniami adresów."""
    rows = [[f"Seller line {i}"] for i in range(18)]
    rows += [["Item No.", "Description", "Qty", "Amount"],
             ["NL753-S-40", "Catheter", "10", "25.00"]]
    doc = extractor.parse_pages([(CI, [rows])])
    assert [i["ref"] for i in doc.items] == ["NL753-S-40"]


def test_description_with_total_is_not_a_total_row():
    table = [["Item No.", "Description", "Qty", "Amount"],
             ["GLV-TC", "Total Care nitrile gloves", "100", "250.00"],
             ["SUMA-1", "Suma vaccum set", "2", "9.00"],
             ["", "Total", "102", "259.00"],
             ["Grand total", "", "", "259.00"]]
    doc = extractor.parse_pages([("COMMERCIAL INVOICE", [table])])
    assert [i["ref"] for i in doc.items] == ["GLV-TC", "SUMA-1"]
    assert doc.total_net == "259.00"


def test_sno_column_numeric_positions_are_not_refs():
    table = [["S.No", "Product", "Qty(pcs)", "Amount"],
             ["1.AT-NFA-S 1", "Catheter", "10", "5.00"],
             ["10", "Mask", "20", "5.00"],
             ["11", "Gloves", "30", "5.00"]]
    doc = extractor.parse_pages([("INVOICE", [table])])
    assert [i["ref"] for i in doc.items] == ["AT-NFA-S_1"]


def test_qty_with_unit_in_one_cell_and_amount_from_qty_times_price():
    table = [["Item No.", "Description", "Quantity", "Unit price"],
             ["NL753-S-40", "Catheter", "1,000 PCS", "0.25"],
             ["MSK-2", "Mask", "20", ""]]
    doc = extractor.parse_pages([("INVOICE", [table])])
    a, b = doc.items
    assert a["qty"] == "1,000" and a["unit"] == "PCS" and a["amount"] == "250"
    assert b["amount"] is None     # brak ceny i wartości → puste, nie cena jednostkowa


def test_invoice_number_ignores_dates():
    assert extractor.extract_invoice_number("Date of invoice: 2026-09-01\nInvoice No.: FV/9") == "FV/9"
    assert extractor.extract_invoice_number("Invoice: 17.09.2026") == ""


def test_blank_page_inside_document_is_continuation(monkeypatch):
    """Rewers duplexu (pusta strona) w środku faktury nie zaczyna części „other”."""
    monkeypatch.setattr(splitter, "page_texts", lambda p: [CI, "", CONT])
    monkeypatch.setattr(splitter, "write_range", lambda *a: None)
    parts = splitter.split_pdf("x/duplex.pdf")
    assert [(p["kind"], p["page_from"], p["page_to"]) for p in parts] == [(InvoiceDocKind.invoice, 1, 3)]


def test_second_invoice_with_distorted_header_starts_new_part(monkeypatch):
    second = "COMMERC1AL 1NVOICE Invoice No.: FV/2 Seller Buyer Item No Qty Amount 123456"
    monkeypatch.setattr(splitter, "page_texts", lambda p: [CI + "\nInvoice No.: FV/1", CONT, second])
    monkeypatch.setattr(splitter, "write_range", lambda *a: None)
    monkeypatch.setattr(splitter.ml, "predict_doc_kind", lambda t: None)
    parts = splitter.split_pdf("x/dwie.pdf")
    assert [(p["kind"], p["page_from"], p["page_to"]) for p in parts] == [
        (InvoiceDocKind.invoice, 1, 2), (InvoiceDocKind.invoice, 3, 3)]
    # ta sama faktura powtórzona w nagłówku strony 2 = kontynuacja
    monkeypatch.setattr(splitter, "page_texts", lambda p: [CI + "\nInvoice No.: FV/1",
                                                            "page 2 Invoice No.: FV/1 " + CONT])
    assert len(splitter.split_pdf("x/jedna.pdf")) == 1


def test_master_header_below_title_rows_and_second_sheet():
    rows = [["Master data ACME", None], ["eksport 2026-09-17"], [],
            ["ref_code", "opis_pl", "Kod CN"], ["R-1", "Opis", "9018"]]
    recs = master_import.parse_workbook_rows(rows)
    assert len(recs) == 1 and recs[0]["tariff_cn"] == "9018"
    # nagłówek jednowierszowy z kolumną „PAZ” nie jest bannerem poziomów
    rows = [["ref_code", "opis_pl", "PAZ"], ["R-2", "Opis", "4800"]]
    assert master_import.parse_workbook_rows(rows)[0]["ref_code"] == "R-2"
    # arkusz „Instrukcja” przed arkuszem z danymi
    import io

    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "Instrukcja"
    wb.active.append(["Jak wypełnić", "..."])
    ws = wb.create_sheet("Dane")
    ws.append(["ref_code", "opis_pl"])
    ws.append(["R-3", "Trzeci"])
    buf = io.BytesIO()
    wb.save(buf)
    rows = master_import.load_xlsx_rows(buf.getvalue())
    assert master_import.parse_workbook_rows(rows)[0]["ref_code"] == "R-3"


def test_conversion_factor_polish_number_format():
    rec = master_import.parse_workbook_rows([
        ["", "", "KAR"], ["ref_code", "Podstawowa jednostka miary", "Ilość podstawowej jednostki miary"],
        ["R-1", "SZT", "1 200"]])[0]
    assert master_import._conversion_rules(rec) == [("CTN", "PCS", 1200.0)]


def test_marker_on_every_page_of_same_invoice_is_one_document(monkeypatch):
    """Wydruk z ERP powtarza „COMMERCIAL INVOICE” + numer na każdej stronie — to jeden
    dokument; inny numer z markerem = nowy dokument."""
    monkeypatch.setattr(splitter, "write_range", lambda *a: None)
    monkeypatch.setattr(splitter.ml, "predict_doc_kind", lambda t: None)
    monkeypatch.setattr(splitter, "page_texts", lambda p: [
        CI + "\nInvoice No.: FV/1", CI + "\nInvoice No.: FV/1 page 2 " + CONT,
        CI + "\nInvoice No.: FV/2"])
    parts = splitter.split_pdf("x/erp.pdf")
    assert [(p["kind"], p["page_from"], p["page_to"]) for p in parts] == [
        (InvoiceDocKind.invoice, 1, 2), (InvoiceDocKind.invoice, 3, 3)]
    # packing lista z numerem tej samej faktury to NIE kontynuacja faktury
    monkeypatch.setattr(splitter, "page_texts", lambda p: [
        CI + "\nInvoice No.: FV/1", PL + "\nInvoice No.: FV/1"])
    assert [p["kind"] for p in splitter.split_pdf("x/cipl.pdf")] == [
        InvoiceDocKind.invoice, InvoiceDocKind.packing_list]


def test_excel_keeps_formula_like_text_as_text(tmp_path):
    """„=HYPERLINK(...)” w REF/opisie z faktury nie może stać się formułą w Excelu."""
    import openpyxl

    from app.invoices.excel import build_workbook
    from app.models import InvoiceItem, InvoiceJob, InvoiceMatchStatus
    item = InvoiceItem(line_no=1, raw_ref='=HYPERLINK("http://evil")', qty="1", amount="=1+1",
                       name_pl="+cmd|' /C calc'!A0", match_status=InvoiceMatchStatus.unmatched,
                       skipped=False)
    job = InvoiceJob(invoice_number="=SUM(A1)", filename="x.pdf", stored_name="x.pdf", items=[item])
    path = tmp_path / "f.xlsx"
    build_workbook([job]).save(path)
    row = list(openpyxl.load_workbook(path).active.iter_rows(min_row=2, max_row=2))[0]
    assert [c.data_type for c in row[:4]] == ["s", "n", "s", "s"]
    assert row[0].value == "=SUM(A1)" and row[2].value == '=HYPERLINK("http://evil")'
    assert row[7].value == "=1+1" and row[7].data_type == "s"


def test_packing_list_weights_are_decimal_cartons_integer():
    out = packing_list.weight_map_from_items([
        {"ref": "A", "weight_net": "1,500", "weight_gross": "1.750", "cartons": "1,000"}])
    assert out["A"] == {"weight_net": "1.500", "weight_gross": "1.750", "cartons": "1000"}


def test_word_rows_fallback_when_ruled_tables_have_no_items():
    """Ramka wokół nagłówka dokumentu = „tabela” bez pozycji; tabela pozycji bez linii —
    wiersze ze słów ratują ekstrakcję. Gdy tabela z linii daje pozycje, fallback nie dubluje."""
    frame = [["Seller", "Buyer"], ["ACME", "ACME"]]
    words = [["Item No.", "Description", "Qty", "Amount"],
             ["NL753-S-40", "Catheter", "1,000", "250.00"]]
    page = extractor.PageData(text="COMMERCIAL INVOICE", tables=[frame], fallback_rows=words)
    doc = extractor.parse_pages([page])
    assert [i["ref"] for i in doc.items] == ["NL753-S-40"]
    ruled = [["Item No.", "Description", "Qty", "Amount"], ["MSK2", "Mask", "20", "200.00"]]
    page = extractor.PageData(text="COMMERCIAL INVOICE", tables=[ruled], fallback_rows=words)
    assert [i["ref"] for i in extractor.parse_pages([page]).items] == ["MSK2"]


def test_visible_overrides_scoped_in_deps():
    from types import SimpleNamespace

    from app.deps import visible_overrides
    from app.models import Role
    material = SimpleNamespace(overrides=[SimpleNamespace(company_id=1), SimpleNamespace(company_id=2)])
    admin = SimpleNamespace(role=Role.admin, company_id=None, view_all_companies=False)
    buyer = SimpleNamespace(role=Role.purchasing, company_id=2, view_all_companies=False)
    assert len(visible_overrides(material, admin)) == 2
    assert [o.company_id for o in visible_overrides(material, buyer)] == [2]


def test_extra_charges_parsed_and_in_excel(tmp_path):
    """Faktury FICTIVA (2026-09-29): „LCL handling charge | US$500.00” bez indeksu między
    pozycjami a „Total” — koszt dodatkowy (suma faktury go zawiera), osobny wiersz w Excelu."""
    import openpyxl

    from app.invoices.excel import build_workbook
    from app.models import InvoiceJob
    table = [["Product Code", "Product Name", "Quantity", "Amount"],
             ["AT-SGS-M 1", "Disposable Gowns", "2500", "US$1,714.50"],
             ["LCL handling charge", "", "", "US$500.00"],
             ["Port of discharge DUBAI", "", "", ""],
             ["Total", "", "2500", "US$2,214.50"]]
    doc = extractor.parse_pages([("COMMERCIAL INVOICE", [table])])
    assert [i["ref"] for i in doc.items] == ["AT-SGS-M 1"]
    assert doc.charges == [{"desc": "LCL handling charge", "amount": "500.00"}]
    job = InvoiceJob(invoice_number="E6399", filename="x.pdf", stored_name="x.pdf", items=[],
                     check_data={"charges": doc.charges})
    path = tmp_path / "f.xlsx"
    build_workbook([job]).save(path)
    row = [c.value for c in list(openpyxl.load_workbook(path).active.iter_rows(min_row=2))[0]]
    assert row[3] == "Koszty dodatkowe: LCL handling charge" and row[7] == 500
